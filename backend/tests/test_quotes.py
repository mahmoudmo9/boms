"""Quotes / proforma invoices: numbering, gate, approval, sending, outcomes, outputs and integrations."""

import io
import threading
from datetime import date, timedelta
from decimal import Decimal
from pathlib import Path

import pytest
from sqlalchemy import create_engine, text
from sqlalchemy.orm import Session

from app.core.db import new_session
from app.modules.parties import clock
from tests.test_costing import day, new_sheet, priced_sheet, selected_case
from tests.test_parties import make_party, tag
from tests.test_sourcing_inquiries import audit_ops, events

# Playwright's headless Chromium build (the full Chromium binary does not print in this container: no D-Bus)
CHROMIUM = next((str(p) for p in Path("/opt/pw-browsers").glob(
    "chromium_headless_shell-*/chrome-headless-shell-linux64/chrome-headless-shell")), "/nonexistent")


def buyer(client, country="AE"):
    return make_party(client, name_en=f"Buyer {tag()}", name_fa=f"خریدار {tag()}", roles=["buyer"],
                      country_iso2=country)


def ready_case(client, country="AE", currencies=("USD", "EUR")):
    """A case with a buyer, two selected offers and a final, priced cost sheet (margin 20 %)."""
    case, items, offers = selected_case(client, currencies)
    party = buyer(client, country)
    client.post("/api/contacts", json={"party_id": party["id"], "name_en": "Mr. Rahimi", "is_primary": True})
    assert client.patch(f"/api/cases/{case['id']}", json={"party_id": party["id"]}).status_code == 200
    sheet = priced_sheet(client, case)
    r = client.post(f"/api/cost-sheets/{sheet['id']}/finalize")
    assert r.status_code == 200, r.text
    return case, items, party, r.json()


def new_quote(client, case, **extra):
    r = client.post("/api/quotes", json={"case_id": case["id"], **extra})
    assert r.status_code == 201, r.text
    return r.json()


def sent_quote(client, **extra):
    case, items, party, sheet = ready_case(client)
    quote = new_quote(client, case, **extra)
    assert client.post(f"/api/quotes/{quote['id']}/approve").status_code == 200
    r = client.post(f"/api/quotes/{quote['id']}/mark-sent", json={})
    assert r.status_code == 200, r.text
    return case, items, party, sheet, r.json()


def interactions_of(client, party_id):
    return client.get("/api/interactions", params={"party_id": party_id}).json()["items"]


# ---- creation, gate, numbering ------------------------------------------------------------------------------

def test_create_from_final_sheet_with_items_terms_and_number(client, db):
    case, items, party, sheet = ready_case(client)
    quote = new_quote(client, case)
    from app.modules.quotes.numbering import numbering_year

    year = numbering_year(date.today(), client.get("/api/me").json()["calendar"])  # the user's calendar decides
    assert quote["number"].startswith(f"PI-{year}-") and quote["version"] == 1 and quote["status"] == "draft"
    assert quote["language"] == "en" and quote["currency"] == "IRR" and quote["price_basis"] == "before_vat"
    assert quote["buyer_party_id"] == party["id"] and quote["contact_name_en"] == "Mr. Rahimi"
    computed = {r["item_id"]: r for r in sheet["computation"]["items"]}
    assert [(i["line_item_id"], i["unit_price"], i["total"]) for i in quote["items"]] == [
        (it["id"], computed[it["id"]]["unit_price"] + ".0000", computed[it["id"]]["total_price"] + ".0000") for it in items]
    assert Decimal(quote["subtotal"]) == sum(Decimal(computed[i]["total_price"]) for i in computed)
    assert quote["valid_until"] == sheet["gate"]["valid_until"]  # default validity capped by FX age / offers
    assert "Dear" in quote["terms_text"] and quote["valid_until"] in quote["terms_text"]
    assert audit_ops(db, "ops.quote", quote["id"])[0] == "I"
    assert events(client, "quote", quote["id"], "create")[0]["data"]["number"] == quote["number"]


def test_gate_blocks_creation_and_lists_failures(client):
    case, items, offers = selected_case(client)
    sheet = new_sheet(client, case)
    r = client.post("/api/quotes", json={"case_id": case["id"]})
    assert r.status_code == 409 and r.json()["error"]["code"] == "gate_failed"
    r = client.post("/api/quotes", json={"case_id": case["id"], "cost_sheet_id": sheet["id"]})
    codes = {f["code"] for f in r.json()["error"]["details"]["failures"]}
    assert {"sheet_not_final", "fx_missing"} <= codes
    assert client.get("/api/quotes", params={"case_id": case["id"]}).json() == []


def test_numbering_is_unique_and_gap_free_under_concurrency(client):
    from app.modules.quotes.numbering import next_number

    year = 3000 + int(tag(), 16) % 5000  # a year nobody else uses
    first = new_session()
    n1 = next_number(first, year)  # holds the year's row until commit
    result = {}

    def other():
        with new_session() as s2:
            result["n2"] = next_number(s2, year)
            s2.commit()

    thread = threading.Thread(target=other)
    thread.start()
    thread.join(timeout=1)
    assert thread.is_alive()  # waits for the first transaction
    first.commit()
    first.close()
    thread.join(timeout=10)
    assert (n1, result["n2"]) == (f"PI-{year}-0001", f"PI-{year}-0002")
    # a rolled-back creation gives its number back
    with new_session() as s3:
        assert next_number(s3, year) == f"PI-{year}-0003"
        s3.rollback()
    with new_session() as s4:
        assert next_number(s4, year) == f"PI-{year}-0003"
        s4.commit()


def test_two_quotes_created_in_parallel_get_consecutive_numbers(client):
    from app.modules.quotes import service

    cases = [ready_case(client)[0] for _ in range(2)]
    barrier = threading.Barrier(2)
    numbers, errors = [], []

    def create(case_id):
        try:
            with new_session() as s:
                barrier.wait()
                quote = service.create_quote(s, {"case_id": case_id}, "gregorian", "script")
                numbers.append(quote.number)
                s.commit()
        except Exception as exc:  # pragma: no cover - reported below
            errors.append(repr(exc))

    threads = [threading.Thread(target=create, args=(c["id"],)) for c in cases]
    for t in threads:
        t.start()
    for t in threads:
        t.join(timeout=30)
    assert not errors, errors
    seqs = sorted(int(n.rsplit("-", 1)[1]) for n in numbers)
    assert len(set(numbers)) == 2 and seqs[1] == seqs[0] + 1


def test_jalali_numbering_year(client):
    from app.modules.quotes.numbering import numbering_year

    assert numbering_year(date(2026, 10, 27), "jalali") == 1405
    assert numbering_year(date(2026, 3, 20), "jalali") == 1404 and numbering_year(date(2026, 3, 21), "jalali") == 1405
    assert numbering_year(date(2026, 10, 27), "gregorian") == 2026


# ---- draft editing, approval lock, immutability -------------------------------------------------------------

def test_draft_edit_manual_price_then_approve_locks(client):
    case, items, party, sheet = ready_case(client)
    quote = new_quote(client, case)
    row = quote["items"][0]
    r = client.patch(f"/api/quotes/items/{row['id']}", json={"unit_price": "99999999", "note": "special price",
                                                             "description_override": "Ball valve DN50 PN16"})
    item = next(i for i in r.json()["items"] if i["id"] == row["id"])
    assert item["price_manual"] is True and Decimal(item["total"]) == Decimal("99999999") * Decimal(item["quantity"])
    r = client.patch(f"/api/quotes/{quote['id']}", json={"price_basis": "with_vat", "lead_time_days": 45,
                                                          "incoterm": "CIF", "incoterm_place": "Bandar Abbas"})
    assert r.status_code == 200
    r = client.patch(f"/api/quotes/{quote['id']}", json={"valid_until": day(60)})
    assert r.status_code == 422 and r.json()["error"]["code"] == "validity_too_long"
    assert client.post(f"/api/quotes/{quote['id']}/mark-sent", json={}).json()["error"]["code"] == "quote_not_approved"
    approved = client.post(f"/api/quotes/{quote['id']}/approve").json()
    assert approved["status"] == "approved" and approved["approved_by"] == "admin" and approved["approved_at"]
    assert client.patch(f"/api/quotes/{quote['id']}", json={"notes": "x"}).json()["error"]["code"] == "quote_locked"
    assert client.patch(f"/api/quotes/items/{row['id']}", json={"note": "x"}).status_code == 409
    assert client.delete(f"/api/quotes/{quote['id']}").status_code == 409
    # a new cost sheet version does not change the approved numbers
    v2 = client.post(f"/api/cost-sheets/{sheet['id']}/new-version").json()
    client.patch(f"/api/cost-sheets/{v2['id']}", json={"target_percent": "35"})
    client.post(f"/api/cost-sheets/{v2['id']}/finalize")
    again = client.get(f"/api/quotes/{quote['id']}").json()
    assert [i["unit_price"] for i in again["items"]] == [i["unit_price"] for i in approved["items"]]


def test_approval_rechecks_minimums(client):
    case, items, party, sheet = ready_case(client)
    # v2: the per-item minimum of the first item equals its computed unit price
    v2 = client.post(f"/api/cost-sheets/{sheet['id']}/new-version").json()
    first = next(r for r in v2["computation"]["items"] if r["item_id"] == items[0]["id"])
    client.patch(f"/api/cost-sheets/{v2['id']}", json={"item_min_prices": {items[0]["id"]: first["unit_price"]}})
    assert client.post(f"/api/cost-sheets/{v2['id']}/finalize").status_code == 200
    quote = new_quote(client, case)
    assert quote["cost_sheet"]["version"] == 2
    client.patch(f"/api/quotes/items/{quote['items'][0]['id']}", json={"unit_price": "1"})
    r = client.post(f"/api/quotes/{quote['id']}/approve")
    assert r.status_code == 409 and r.json()["error"]["code"] == "gate_failed"
    assert r.json()["error"]["details"]["failures"] == [{"code": "item_below_minimum", "item_ids": [items[0]["id"]]}]
    # back to the sheet's price: approvable
    client.patch(f"/api/quotes/items/{quote['items'][0]['id']}", json={"unit_price": None})
    assert client.post(f"/api/quotes/{quote['id']}/approve").status_code == 200


def test_mark_sent_one_interaction_followup_and_funnel(client, db):
    case, items, party, sheet, quote = sent_quote(client)
    assert quote["status"] == "sent" and quote["sent_at"]
    rows = interactions_of(client, party["id"])
    assert len(rows) == 1
    detail = client.get(f"/api/interactions/{rows[0]['id']}").json()
    assert detail["direction"] == "out" and quote["number"] in detail["subject"]
    fu = detail["followups"][0]
    assert fu["action"] == "Chase quote response"
    assert fu["due_on"] == max(date.today(), date.fromisoformat(quote["valid_until"]) - timedelta(days=1)).isoformat()
    ev = events(client, "quote", quote["id"], "sent")[0]["data"]
    assert ev["followup_id"] == fu["id"] and ev["funnel_stage"] == "quoted"
    roles = {r["role"]: r for r in client.get(f"/api/parties/{party['id']}").json()["roles"]}
    assert roles["buyer"]["stage"] == "quoted"
    # sent is immutable
    assert client.patch(f"/api/quotes/{quote['id']}", json={"notes": "x"}).status_code == 409
    assert client.post(f"/api/quotes/{quote['id']}/mark-sent", json={}).status_code == 409


def test_mark_sent_without_funnel_move_and_with_no_followup_reason(client):
    case, items, party, sheet = ready_case(client)
    quote = new_quote(client, case)
    client.post(f"/api/quotes/{quote['id']}/approve")
    r = client.post(f"/api/quotes/{quote['id']}/mark-sent", json={"move_to_quoted": False, "channel": "whatsapp",
                                                                   "no_followup_reason": "buyer will call"})
    assert r.status_code == 200
    roles = {x["role"]: x for x in client.get(f"/api/parties/{party['id']}").json()["roles"]}
    assert roles["buyer"]["stage"] != "quoted"
    ev = events(client, "quote", quote["id"], "sent")[0]["data"]
    assert ev["followup_id"] is None and ev["no_followup_reason"] == "buyer will call"


def test_revise_creates_new_version_and_supersedes(client, db):
    case, items, party, sheet, quote = sent_quote(client)
    client.patch  # noqa: B018
    r = client.post(f"/api/quotes/{quote['id']}/revise", json={})
    assert r.status_code == 201, r.text
    new = r.json()
    assert new["version"] == 2 and new["status"] == "draft" and new["supersedes_id"] == quote["id"]
    assert new["number"] != quote["number"]
    assert client.get(f"/api/quotes/{quote['id']}").json()["status"] == "superseded"
    fu_status = db.execute(text("SELECT status, done_note FROM core.followup WHERE interaction_id = :i"),
                           {"i": quote["interaction_id"]}).one()
    assert tuple(fu_status) == ("cancelled", "quote revised")
    listed = client.get(f"/api/cases/{case['id']}/quotes").json()
    assert [q["version"] for q in listed["items"]] == [2, 1]


def test_expiry_job_changes_status_only(client, db, monkeypatch):
    case, items, party, sheet, quote = sent_quote(client)
    before = db.execute(text("SELECT count(*) FROM core.interaction")).scalar()
    from datetime import UTC, datetime

    later = datetime.combine(date.fromisoformat(quote["valid_until"]) + timedelta(days=1), datetime.min.time(),
                             tzinfo=UTC) + timedelta(hours=12)
    monkeypatch.setattr(clock, "now", lambda: later)
    monkeypatch.setattr(clock, "local_tz", lambda: UTC)
    from app.modules.quotes.handlers import expire

    assert expire(db, {})["quotes_expired"] >= 1
    db.commit()
    assert client.get(f"/api/quotes/{quote['id']}").json()["status"] == "expired"
    assert db.execute(text("SELECT count(*) FROM core.interaction")).scalar() == before
    assert events(client, "quote", quote["id"], "status")[0]["data"]["new_status"] == "expired"


# ---- outcomes ---------------------------------------------------------------------------------------------

def test_outcomes_derive_status_and_drive_sold(client, db):
    case, items, party, sheet, quote = sent_quote(client)
    a, b = quote["items"]
    r = client.put(f"/api/quotes/{quote['id']}/outcomes", json={"items": [
        {"quote_item_id": b["id"], "outcome": "rejected"}]})
    assert r.status_code == 422 and r.json()["error"]["code"] == "reason_required"
    r = client.put(f"/api/quotes/{quote['id']}/outcomes", json={"items": [
        {"quote_item_id": a["id"], "outcome": "accepted"},
        {"quote_item_id": b["id"], "outcome": "rejected", "outcome_reason": "price", "competitor_price": "123",
         "competitor_currency": "USD", "outcome_note": "competitor 5 % lower"}]})
    assert r.status_code == 200 and r.json()["status"] == "partially_accepted"
    item = lambda iid: client.get(f"/api/line-items/{iid}").json()  # noqa: E731
    assert item(a["line_item_id"])["status"] == "sold" and item(b["line_item_id"])["status"] == "selected"
    # chase follow-up closed when decided
    assert db.execute(text("SELECT status FROM core.followup WHERE interaction_id = :i"),
                      {"i": quote["interaction_id"]}).scalar() == "cancelled"
    # an offer change does not touch a sold item
    offer_id = client.get(f"/api/line-items/{a['line_item_id']}").json()["selected_offer_item_id"]
    offer = db.execute(text("SELECT offer_id FROM ops.offer_item WHERE id = :i"), {"i": offer_id}).scalar()
    client.post(f"/api/offers/{offer}/status", json={"status": "withdrawn"})
    assert item(a["line_item_id"])["status"] == "sold"
    # all accepted -> accepted; revising the outcome clears sold
    r = client.put(f"/api/quotes/{quote['id']}/outcomes", json={"items": [
        {"quote_item_id": b["id"], "outcome": "accepted"}]})
    assert r.json()["status"] == "accepted" and item(b["line_item_id"])["status"] == "sold"
    r = client.put(f"/api/quotes/{quote['id']}/outcomes", json={"items": [
        {"quote_item_id": a["id"], "outcome": "rejected", "outcome_reason": "lead_time"},
        {"quote_item_id": b["id"], "outcome": "rejected", "outcome_reason": "no_budget"}]})
    assert r.json()["status"] == "rejected"
    assert item(a["line_item_id"])["status"] == "uncovered" or item(a["line_item_id"])["status"] != "sold"
    assert item(b["line_item_id"])["status"] != "sold"
    ev = [e["data"] for e in events(client, "line_item", b["line_item_id"], "status")]
    assert any(e["new_status"] == "sold" for e in ev) and any(e["old_status"] == "sold" for e in ev)


def test_delivered_and_closed_items_are_never_touched(client, db):
    case, items, party, sheet, quote = sent_quote(client)
    a, b = quote["items"]
    db.execute(text("UPDATE ops.line_item SET status = 'delivered' WHERE id = :i"), {"i": a["line_item_id"]})
    db.commit()
    client.put(f"/api/quotes/{quote['id']}/outcomes", json={"items": [
        {"quote_item_id": a["id"], "outcome": "accepted"}, {"quote_item_id": b["id"], "outcome": "accepted"}]})
    assert client.get(f"/api/line-items/{a['line_item_id']}").json()["status"] == "delivered"
    assert client.get(f"/api/line-items/{b['line_item_id']}").json()["status"] == "sold"


def test_case_result_once_quotes_are_decided(client):
    case, items, party, sheet, quote = sent_quote(client)
    r = client.post(f"/api/cases/{case['id']}/quote-result", json={"result": "won"})
    assert r.status_code == 409 and r.json()["error"]["code"] == "quotes_undecided"
    client.put(f"/api/quotes/{quote['id']}/outcomes", json={"items": [
        {"quote_item_id": i["id"], "outcome": "accepted"} for i in quote["items"]]})
    r = client.post(f"/api/cases/{case['id']}/quote-result", json={"result": "won"})
    assert r.status_code == 200 and r.json()["won"] == 1
    assert client.get(f"/api/cases/{case['id']}").json()["status"] == "won"


def test_withdraw_and_delete_rules(client):
    case, items, party, sheet = ready_case(client)
    draft = new_quote(client, case)
    assert client.post(f"/api/quotes/{draft['id']}/withdraw", json={"reason": "x"}).status_code == 409
    assert client.delete(f"/api/quotes/{draft['id']}").status_code == 204
    second = new_quote(client, case)
    client.post(f"/api/quotes/{second['id']}/approve")
    r = client.post(f"/api/quotes/{second['id']}/withdraw", json={"reason": "buyer cancelled the project"})
    assert r.json()["status"] == "withdrawn"
    # the sheet a quote uses cannot be deleted (it is final anyway); drafts used by quotes neither
    assert client.delete(f"/api/cost-sheets/{sheet['id']}").status_code == 409


# ---- outputs ------------------------------------------------------------------------------------------------

def _setup_letterhead(client):
    client.patch("/api/settings", json={
        "company.name_en": "Kimia Trading Co.", "company.name_fa": "شرکت بازرگانی کیمیا",
        "company.address": {"en": "No. 5, Example St., Tehran", "fa": "تهران، خیابان نمونه، پلاک ۵"},
        "company.phone": "+98 21 1234 5678", "company.email": "sales@kimia.example",
        "company.bank_details": {"en": "Bank Melli, IBAN IR00 0000", "fa": "بانک ملی، شبا IR00 0000"},
        "company.footer": {"en": "Thank you for your business", "fa": "با سپاس از همکاری شما"}})


def test_print_html_en_and_fa_with_draft_watermark(client):
    _setup_letterhead(client)
    case, items, party, sheet = ready_case(client)
    quote = new_quote(client, case)
    page = client.get(f"/api/quotes/{quote['id']}/print").text
    for needle in ("Proforma invoice", quote["number"], "Kimia Trading Co.", "Valve", "Pump", "DRAFT",
                   "All prices are before VAT.", "Mr. Rahimi", 'dir="ltr"', "@font-face", "Thank you for your business"):
        assert needle in page, needle
    client.patch(f"/api/quotes/{quote['id']}", json={"language": "fa", "price_basis": "with_vat"})
    client.patch("/api/me", json={"calendar": "jalali", "digits": "fa"})
    try:
        page = client.get(f"/api/quotes/{quote['id']}/print").text
    finally:
        client.patch("/api/me", json={"calendar": "gregorian", "digits": "latin"})
    for needle in ("پیش‌فاکتور", 'dir="rtl"', "شرکت بازرگانی کیمیا", "مالیات بر ارزش افزوده", "۱۴۰۵", "پیش‌نویس"):
        assert needle in page, needle
    client.post(f"/api/quotes/{quote['id']}/approve")
    page = client.get(f"/api/quotes/{quote['id']}/print").text
    assert "class='watermark'" not in page and 'class="watermark"' not in page


def test_word_output_reopens_with_items_totals_and_is_linked(client):
    from docx import Document

    _setup_letterhead(client)
    case, items, party, sheet = ready_case(client)
    quote = new_quote(client, case)
    r = client.post(f"/api/quotes/{quote['id']}/outputs", json={"format": "docx"})
    assert r.status_code == 201, r.text
    out = r.json()
    assert out["draft"] is True and out["filename"] == f"{quote['number']}-v1-DRAFT.docx"
    data = client.get(f"/api/documents/{out['document_id']}/download").content
    doc = Document(io.BytesIO(data))
    body = "\n".join(p.text for p in doc.paragraphs)
    cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    header = "\n".join(p.text for p in doc.sections[0].header.paragraphs)
    assert "Proforma invoice" in body and quote["number"] in body and "DRAFT" in body + header
    assert any("Valve" in c for c in cells) and any("Pump" in c for c in cells) and any("Subtotal" in c for c in cells)
    assert "Kimia Trading Co." in header
    links = client.get(f"/api/quotes/{quote['id']}").json()["outputs"]
    assert [link["document_id"] for link in links] == [out["document_id"]]
    # Persian: RTL paragraphs
    client.patch(f"/api/quotes/{quote['id']}", json={"language": "fa"})
    fa = client.post(f"/api/quotes/{quote['id']}/outputs", json={"format": "docx"}).json()
    fa_doc = Document(io.BytesIO(client.get(f"/api/documents/{fa['document_id']}/download").content))
    assert "پیش‌فاکتور" in "\n".join(p.text for p in fa_doc.paragraphs)
    assert fa_doc.paragraphs[0]._p.pPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}bidi") is not None


def test_pdf_unavailable_with_invalid_browser_path(client):
    case, items, party, sheet = ready_case(client)
    quote = new_quote(client, case)
    client.patch("/api/settings", json={"pdf.browser_path": "/nonexistent/msedge.exe"})
    try:
        r = client.post(f"/api/quotes/{quote['id']}/outputs", json={"format": "pdf"})
    finally:
        client.patch("/api/settings", json={"pdf.browser_path": None})
    assert r.status_code == 424 and r.json()["error"]["code"] == "pdf_unavailable"
    assert r.json()["error"]["details"]["reason"] == "browser_path_invalid"
    assert client.get(f"/api/quotes/{quote['id']}").json()["outputs"] == []


@pytest.mark.skipif(not Path(CHROMIUM).exists(), reason="no Chromium-family browser in this test environment")
def test_pdf_with_the_available_chromium(client):
    _setup_letterhead(client)
    case, items, party, sheet = ready_case(client)
    quote = new_quote(client, case)
    client.post(f"/api/quotes/{quote['id']}/approve")
    client.patch("/api/settings", json={"pdf.browser_path": CHROMIUM})
    try:
        r = client.post(f"/api/quotes/{quote['id']}/outputs", json={"format": "pdf"})
    finally:
        client.patch("/api/settings", json={"pdf.browser_path": None})
    assert r.status_code == 201, r.text
    out = r.json()
    assert out["draft"] is False and out["filename"] == f"{quote['number']}-v1.pdf"
    data = client.get(f"/api/documents/{out['document_id']}/download").content
    assert data[:5] == b"%PDF-"
    import fitz  # pymupdf (already a dependency): the PDF text has the quote number and an item

    with fitz.open(stream=data, filetype="pdf") as pdf:
        content = "".join(page.get_text() for page in pdf)
    assert quote["number"] in content and "Valve" in content and "DRAFT" not in content


# ---- integrations -------------------------------------------------------------------------------------------

def test_product_page_quotes_we_sent(client):
    case, items, party, sheet, quote = sent_quote(client)
    made = client.post(f"/api/line-items/{items[0]['id']}/create-product", json={"confirm_duplicate": True}).json()
    data = client.get(f"/api/quotes/by-product/{made['product_id']}").json()
    assert [q["number"] for q in data["quotes"]] == [quote["number"]]
    row = data["quotes"][0]
    assert row["party_id"] == party["id"] and row["outcome"] is None and row["currency"] == "IRR"
    summary = data["summary"][0]
    assert summary["count"] == 1 and summary["min"] == summary["median"] == summary["max"]


def test_party_delete_counts_quotes(client):
    case, items, party, sheet, quote = sent_quote(client)
    r = client.delete(f"/api/parties/{party['id']}")
    assert r.status_code == 409 and r.json()["error"]["details"]["quotes"] == 1
    r = client.request("DELETE", f"/api/parties/{party['id']}?cascade=true", json={"confirm_name": party["name_en"]})
    assert r.status_code == 409 and r.json()["error"]["details"]["quotes"] == 1


def test_case_cascade_delete_removes_sheets_quotes_and_cancels_chases(client, db):
    case, items, party, sheet, quote = sent_quote(client)
    out = client.post(f"/api/quotes/{quote['id']}/outputs", json={"format": "docx"}).json()
    r = client.request("DELETE", f"/api/cases/{case['id']}?cascade=true", json={"confirm_title": case["title"]})
    assert r.status_code == 200, r.text
    counts = r.json()["counts"]
    assert counts["quotes"] == 1 and counts["cost_sheets"] == 1 and counts["quote_outputs"] == 1
    assert counts["followups_cancelled"] >= 1
    assert client.get(f"/api/quotes/{quote['id']}").status_code == 404
    assert client.get(f"/api/cost-sheets/{sheet['id']}").status_code == 404
    assert client.get(f"/api/documents/{out['document_id']}").status_code == 200  # files stay
    assert db.execute(text("SELECT status, done_note FROM core.followup WHERE interaction_id = :i"),
                      {"i": quote["interaction_id"]}).one() == ("cancelled", "case deleted")
    assert audit_ops(db, "ops.quote", quote["id"])[-1] == "D"
    assert audit_ops(db, "ops.cost_sheet", sheet["id"])[-1] == "D"
    snap = events(client, "case", case["id"], "delete")[0]["data"]["deleted"]
    assert snap["quotes"][0]["quote"]["number"] == quote["number"] and snap["cost_sheets"]


def test_search_finds_quotes_by_number_and_buyer(client):
    case, items, party, sheet = ready_case(client)
    quote = new_quote(client, case)
    for q in (quote["number"], party["name_en"]):
        groups = {g["type"]: [h["id"] for h in g["items"]]
                  for g in client.get("/api/search", params={"q": q, "limit": 20}).json()["groups"]}
        assert quote["id"] in groups.get("quote", []), q


def test_request_models_reject_unknown_fields(client):
    for method, path, body in [
        ("post", "/api/quotes", {"case_id": "00000000-0000-0000-0000-000000000000", "bogus": 1}),
        ("patch", "/api/quotes/00000000-0000-0000-0000-000000000000", {"bogus": 1}),
        ("post", "/api/quotes/00000000-0000-0000-0000-000000000000/mark-sent", {"bogus": 1}),
        ("put", "/api/quotes/00000000-0000-0000-0000-000000000000/outcomes", {"items": [], "bogus": 1}),
        ("post", "/api/tax-parameters", {"kind": "vat", "bogus": 1}),
        ("post", "/api/cost-sheets", {"bogus": 1}),
        ("put", "/api/cost-sheets/00000000-0000-0000-0000-000000000000/fx", {"bogus": 1}),
        ("post", "/api/cost-templates", {"bogus": 1}),
    ]:
        r = getattr(client, method)(path, json=body)
        assert r.status_code == 422, (path, r.status_code)


def test_constraint_errors_return_stable_codes_and_persist_nothing(client, db):
    case, items, offers = selected_case(client)
    sheet = new_sheet(client, case)
    before = db.execute(text("SELECT count(*) FROM ops.fx_rate_entry")).scalar()
    r = client.put(f"/api/cost-sheets/{sheet['id']}/fx", json={"from_currency": "IRR", "to_currency": "IRR",
                                                               "rate": "1", "as_of": day()})
    assert r.status_code == 422 and r.json()["error"]["code"] == "invalid_fx"
    assert "fx_rate_entry" not in r.json()["error"]["message"] and "constraint" not in r.json()["error"]["message"]
    assert db.execute(text("SELECT count(*) FROM ops.fx_rate_entry")).scalar() == before
    r = client.post("/api/tax-parameters", json={"kind": "vat", "name": "VAT", "rate_percent": "10",
                                                 "effective_from": day(), "confirmed_by": "x"})
    assert r.status_code == 422 and r.json()["error"]["code"] == "confirmation_incomplete"


def test_terms_templates_reject_unknown_placeholders(client):
    r = client.post("/api/message-templates", json={"kind": "quote_terms", "language": "en", "channel": "letter",
                                                    "name": f"Terms {tag()}", "body_template": "Hi {supplier_name}"})
    assert r.status_code == 422 and r.json()["error"]["details"]["unknown"] == ["supplier_name"]
    r = client.post("/api/message-templates", json={"kind": "quote_terms", "language": "en", "channel": "letter",
                                                    "name": f"Terms {tag()}", "body_template": "Dear {buyer_name}, {bank_details}"})
    assert r.status_code == 201
    listed = client.get("/api/message-templates", params={"kind": "quote_terms"}).json()
    assert "bank_details" in listed["placeholders"] and all(t["kind"] == "quote_terms" for t in listed["items"])
    assert all(t["kind"] == "rfq" for t in client.get("/api/message-templates").json()["items"])


def test_command_center_sales_numbers_on_a_fixed_dataset():
    """Fresh database: one draft awaiting approval, one sent quote expiring, one stale-FX sheet, one won, one lost."""
    from alembic import command
    from fastapi.testclient import TestClient

    from app.cli import alembic_config
    from app.core import dbadmin
    from app.core.config import get_config
    from app.modules.parties import command_center
    from tests.conftest import create_empty_database, superuser_conninfo

    name = "boms_test_cc_quotes"
    create_empty_database(name)
    url = get_config().db_url(name)
    command.upgrade(alembic_config(url), "head")
    engine = create_engine(url)
    try:
        from app.modules.costing import service as costing
        from app.modules.parties import service as parties
        from app.modules.quotes import service as quotes
        from app.modules.sourcing import coverage, inquiries, items as items_service, offers
        from app.modules.system.seed import seed_base_data

        today = clock.today()
        with Session(engine) as db:
            seed_base_data(db)
            sup, _ = parties.create_party(db, {"name_en": "Supplier", "roles": ["supplier"], "confirm_duplicate": True}, "script")
            buy, _ = parties.create_party(db, {"name_en": "Buyer", "roles": ["buyer"], "confirm_duplicate": True}, "script")

            def ready(title, fx_age=0):
                case = inquiries.create_case(db, {"title": title, "party_id": buy.id}, "script")
                (item,) = items_service.add_items(db, None, case.id, [{"name": "X", "quantity": Decimal(2)}], "script")
                off = offers.record_offer(db, {"case_id": case.id, "party_id": sup.id, "channel": "email",
                                               "currency": "USD", "valid_until": today + timedelta(days=30),
                                               "items": [{"line_item_id": item.id, "unit_price": Decimal(10)}],
                                               "no_followup_reason": "fyi"}, "script")
                coverage.select_offer(db, item.id, offers.offer_items(db, off.id)[0].id, None, "script")
                sheet = costing.create_sheet(db, {"case_id": case.id, "deal_type": "domestic",
                                                  "pricing_currency": "USD", "blank": True}, "script")
                costing.update_sheet(db, sheet.id, {"target_percent": Decimal(20)}, "script")
                if fx_age:
                    costing.set_fx(db, sheet.id, {"from_currency": "EUR", "to_currency": "USD", "rate": Decimal("1.1"),
                                                  "as_of": today - timedelta(days=fx_age)}, "script")
                costing.finalize(db, sheet.id, "script")
                return case

            draft_case = ready("Draft")
            quotes.create_quote(db, {"case_id": draft_case.id}, "gregorian", "script")

            def sent(title, valid_days):
                case = ready(title)
                q = quotes.create_quote(db, {"case_id": case.id, "valid_until": today + timedelta(days=valid_days)},
                                        "gregorian", "script")
                quotes.approve(db, q.id, "script", "owner")
                quotes.mark_sent(db, q.id, {"no_followup_reason": "x"}, "script", "owner")
                return q

            passed = sent("Passed", 1)
            quotes.expire_quotes(db, today + timedelta(days=2))  # past its validity, no outcome: still listed
            expiring = sent("Expiring", 2)
            sent("Later", 20)
            won = sent("Won", 20)
            quotes.set_outcomes(db, won.id, [{"quote_item_id": i.id, "outcome": "accepted"}
                                             for i in quotes.items_of(db, won.id)], "script", "owner")
            lost = sent("Lost", 20)
            quotes.set_outcomes(db, lost.id, [{"quote_item_id": i.id, "outcome": "rejected", "outcome_reason": "price"}
                                              for i in quotes.items_of(db, lost.id)], "script", "owner")
            stale = ready("Stale FX", fx_age=10)
            db.commit()
            section = command_center.summary(db)["sourcing"]
            assert [q["case_title"] for q in section["quotes_awaiting_approval"]] == ["Draft"]
            assert [q["quote_id"] for q in section["quotes_expiring"]] == [str(passed.id), str(expiring.id)]
            assert [q["days_left"] for q in section["quotes_expiring"]] == [1, 2]
            assert [s["case_title"] for s in section["sheets_stale_fx"]] == ["Stale FX"]
            assert (section["quotes_won_30_days"], section["quotes_lost_30_days"]) == (1, 1)
            assert str(stale.id)
    finally:
        engine.dispose()
        dbadmin.drop_database(superuser_conninfo(), name)


def test_migrations_0007_0008_up_down_up_with_phase5_data():
    from alembic import command

    from app.cli import alembic_config
    from app.core import dbadmin
    from app.core.config import get_config
    from tests.conftest import create_empty_database, superuser_conninfo

    name = "boms_test_p6_migrations"
    create_empty_database(name)
    url = get_config().db_url(name)
    cfg = alembic_config(url)
    engine = create_engine(url)
    new_tables = {"ops.tax_kind", "ops.tax_parameter", "ops.cost_template", "ops.cost_template_line", "ops.cost_sheet",
                  "ops.fx_rate_entry", "ops.cost_line", "ops.outcome_reason", "ops.quote_sequence", "ops.quote",
                  "ops.quote_item"}

    def tables():
        with engine.connect() as conn:
            return set(conn.execute(text("SELECT schemaname || '.' || tablename FROM pg_tables "
                                         "WHERE schemaname = 'ops'")).scalars())

    def phase5_rows():
        with engine.connect() as conn:
            return conn.execute(text(
                "SELECT (SELECT count(*) FROM ops.\"case\"), (SELECT count(*) FROM ops.line_item), "
                "(SELECT count(*) FROM ops.message_template WHERE kind = 'rfq'), (SELECT count(*) FROM core.party)")).one()

    try:
        command.upgrade(cfg, "0006")
        with engine.begin() as conn:  # Phase 5 data as it exists before the update
            party = conn.execute(text("INSERT INTO core.party (name_en) VALUES ('Old buyer') RETURNING id")).scalar()
            case = conn.execute(text("INSERT INTO ops.\"case\" (title, party_id) VALUES ('Old case', :p) RETURNING id"),
                                {"p": party}).scalar()
            conn.execute(text("INSERT INTO ops.line_item (case_id, name, quantity) VALUES (:c, 'Old item', 3)"),
                         {"c": case})
        before = phase5_rows()
        assert not (new_tables & tables())
        command.upgrade(cfg, "head")
        assert new_tables <= tables() and phase5_rows() == before
        with engine.connect() as conn:
            assert conn.execute(text("SELECT count(*) FROM ops.cost_template WHERE is_default")).scalar() == 3
            assert conn.execute(text("SELECT count(*) FROM ops.tax_parameter")).scalar() == 0  # no rates seeded
            assert conn.execute(text("SELECT count(*) FROM ops.tax_kind")).scalar() == 5
            assert conn.execute(text("SELECT count(*) FROM ops.message_template WHERE kind = 'quote_terms'")).scalar() == 2
        command.downgrade(cfg, "0006")
        assert not (new_tables & tables()) and phase5_rows() == before
        command.upgrade(cfg, "head")
        assert new_tables <= tables() and phase5_rows() == before
    finally:
        engine.dispose()
        dbadmin.drop_database(superuser_conninfo(), name)


def test_backup_and_restore_test_cover_every_table_with_quote_data(client):
    from app.core.db import get_engine
    from app.modules.system import backup

    sent_quote(client)
    run_id = backup.run_backup("manual", source="script")
    result = backup.restore_test(run_id, source="script")
    with get_engine().connect() as conn:
        counts = backup.table_counts(conn)
    assert result["status"] == "passed", result
    assert result["details"]["tables_checked"] == len(counts) and result["details"]["row_count_mismatches"] == {}
    for table in ("ops.cost_sheet", "ops.cost_line", "ops.fx_rate_entry", "ops.quote", "ops.quote_item",
                  "ops.quote_sequence", "ops.tax_kind", "ops.outcome_reason"):
        assert counts[table] > 0, table


def test_audit_rows_and_events_for_every_new_table(client, db):
    case, items, party, sheet, quote = sent_quote(client)
    client.put(f"/api/quotes/{quote['id']}/outcomes", json={"items": [
        {"quote_item_id": i["id"], "outcome": "accepted"} for i in quote["items"]]})
    client.post("/api/tax-parameters", json={"kind": "vat", "name": f"VAT {tag()}", "rate_percent": "10",
                                             "effective_from": "2020-01-01", "effective_to": "2020-01-02"})
    client.post("/api/cost-templates", json={"name": f"T {tag()}", "deal_type": "domestic", "lines": [
        {"kind": "purchase", "label_en": "Purchase", "basis": "per_unit"}]})
    audited = set(db.execute(text("SELECT DISTINCT table_name FROM core.audit")).scalars())
    for table in ("ops.tax_parameter", "ops.cost_template", "ops.cost_template_line", "ops.cost_sheet",
                  "ops.fx_rate_entry", "ops.cost_line", "ops.quote", "ops.quote_item", "ops.quote_sequence"):
        assert table in audited, table
    entity_types = set(db.execute(text("SELECT DISTINCT entity_type FROM core.event")).scalars())
    assert {"tax_parameter", "cost_template", "cost_sheet", "quote"} <= entity_types
    actions = {e["action"] for e in client.get("/api/events", params={"entity_type": "quote", "entity_id": quote["id"],
                                                                        "page_size": 50}).json()["items"]}
    assert {"create", "approve", "sent", "outcome"} <= actions
