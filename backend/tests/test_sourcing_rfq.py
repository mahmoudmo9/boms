"""Phase 5 milestone 2: RFQs, templates, offers, coverage, comparison, catalog, capabilities, command center,
search, expiry handlers."""

from datetime import UTC, date, datetime, timedelta
from decimal import Decimal

import pytest
from sqlalchemy import create_engine, func, select, text
from sqlalchemy.orm import Session

from app.cli import alembic_config
from app.core import dbadmin
from app.core.config import get_config
from app.modules.parties import clock
from app.modules.sourcing import coverage, dashboard, offers, rfq
from app.modules.sourcing.models import Rfq
from tests.conftest import create_empty_database, superuser_conninfo, unique_bytes, upload
from tests.test_parties import make_party, tag
from tests.test_sourcing_inquiries import audit_ops, events

TODAY = date.today()


def soon(days: int = 7) -> str:
    return (TODAY + timedelta(days=days)).isoformat()


def make_case(client, items, **extra):
    case = client.post("/api/cases", json={"title": f"Case {tag()}", **extra}).json()
    r = client.post("/api/line-items/bulk", json={"case_id": case["id"], "items": items})
    assert r.status_code == 201, r.text
    return case, r.json()["items"]


def supplier(client, name=None, country=None, **patch):
    party = make_party(client, name_en=name or f"Supplier {tag()}", roles=["supplier"], country_iso2=country)
    if patch:
        party = client.patch(f"/api/parties/{party['id']}", json=patch).json()
    return party


def contact(client, party_id, email=None, name="Ali"):
    methods = [{"type": "email", "value": email}] if email else []
    r = client.post("/api/contacts", json={"party_id": party_id, "name_en": name, "methods": methods})
    assert r.status_code == 201, r.text
    return r.json()


def new_rfq(client, case, items, recipients, **extra):
    body = {"case_id": case["id"], "respond_by": soon(5), "items": [{"line_item_id": i["id"]} for i in items],
            "recipients": recipients, **extra}
    r = client.post("/api/rfqs", json=body)
    assert r.status_code == 201, r.text
    return r.json()


def offer_body(case, party, rows, **extra):
    return {"case_id": case["id"], "party_id": party["id"], "channel": "email", "currency": "USD",
            "valid_until": soon(30), "items": rows, "no_followup_reason": "comparing offers", **extra}


def interactions_of(client, party_id):
    return client.get("/api/interactions", params={"party_id": party_id}).json()["items"]


# ---- templates and rendering ------------------------------------------------------------------

def test_render_with_all_placeholders_en_and_fa(client):
    client.patch("/api/settings", json={"company.name": "Kimia Trading",
                                        "company.signature": {"en": "Best regards,\nKimia Trading", "fa": "با احترام\nبازرگانی کیمیا"}})
    case, items = make_case(client, [
        {"name": "Polyethylene HD", "quantity": "20", "unit_code": "t", "packaging": "25 kg bags",
         "incoterm": "CIF", "incoterm_place": "Bandar Abbas", "destination_country_iso2": "IR",
         "required_by": "2026-12-01", "specs": [{"key": "Grade", "value": "PE100"}, {"key": "MFI", "value": "0.3"}]},
        {"name": "Sulphuric acid", "quantity": "5000", "unit_code": "kg"}])
    foreign = supplier(client, name=f"Sinochem {tag()}", country="CN")
    person = contact(client, foreign["id"], email=f"sales{tag()}@sino.example", name="Wang Lei")
    domestic = make_party(client, name_fa=f"پتروشیمی نمونه {tag()}", roles=["supplier"], country_iso2="IR")
    body = new_rfq(client, case, items, [{"party_id": foreign["id"], "contact_id": person["id"]},
                                         {"party_id": domestic["id"]}], respond_by="2026-11-20")
    en, fa = body["recipients"]
    assert (en["channel"], en["language"], fa["channel"], fa["language"]) == ("email", "en", "whatsapp", "fa")
    assert en["subject"] == f"Request for quotation: {case['title']}"
    text_en = en["message_text"]
    assert text_en.startswith("Dear Wang Lei,")
    assert "1. Polyethylene HD — Quantity: 20 Metric ton" in text_en
    assert "Specs: Grade: PE100; MFI: 0.3" in text_en and "Packaging: 25 kg bags" in text_en
    assert "Incoterm: CIF Bandar Abbas" in text_en and "Destination: Iran" in text_en
    assert "Required by: 2026-12-01" in text_en and "2. Sulphuric acid — Quantity: 5000 Kilogram" in text_en
    assert "by 2026-11-20" in text_en and text_en.rstrip().endswith("Best regards,\nKimia Trading")
    text_fa = fa["message_text"]
    assert text_fa.startswith(f"سلام {domestic['name_fa']}،")  # no contact: the supplier's name
    assert "مقدار: 20 تن" in text_fa and "مشخصات: Grade: PE100" in text_fa and "تا 1405/08/29" in text_fa
    assert text_fa.rstrip().endswith("با احترام\nبازرگانی کیمیا")
    assert "{" not in text_en + text_fa
    assert en["mailto"].startswith(f"mailto:{person['methods'][0]['value']}?subject=Request%20for%20quotation")
    assert fa["mailto"] is None


def test_unknown_placeholder_rejected(client):
    base = {"language": "en", "channel": "email", "name": f"T {tag()}", "body_template": "Hi {contact_name}"}
    r = client.post("/api/message-templates", json={**base, "body_template": "Hi {contact} {items_table}"})
    assert r.status_code == 422 and r.json()["error"]["details"]["unknown"] == ["contact"]
    r = client.post("/api/message-templates", json={**base, "subject_template": "Re {price}"})
    assert r.json()["error"]["code"] == "unknown_placeholder"
    assert client.post("/api/message-templates", json={**base, "body_template": "Hi {"}).json()["error"]["code"] == \
        "invalid_template"
    created = client.post("/api/message-templates", json={**base, "body_template": "Use {{braces}} {signature}"})
    assert created.status_code == 201
    r = client.patch(f"/api/message-templates/{created.json()['id']}", json={"body_template": "{unknown}"})
    assert r.status_code == 422
    preview = client.post("/api/message-templates/preview", json={"body_template": "Use {{braces}} for {supplier_name}"})
    assert preview.json()["body"] == "Use {braces} for ACME Trading"


def test_template_default_per_language_channel_and_delete_rules(client):
    t1 = client.post("/api/message-templates", json={"language": "de", "channel": "email", "name": "DE 1",
                                                     "body_template": "Hallo {contact_name}", "is_default": True}).json()
    t2 = client.post("/api/message-templates", json={"language": "de", "channel": "email", "name": "DE 2",
                                                     "body_template": "Guten Tag {contact_name}", "is_default": True}).json()
    rows = {t["id"]: t for t in client.get("/api/message-templates").json()["items"]}
    assert rows[t2["id"]]["is_default"] and not rows[t1["id"]]["is_default"]
    case, items = make_case(client, [{"name": "Gear"}])
    party = supplier(client, country="DE")
    body = new_rfq(client, case, items, [{"party_id": party["id"], "language": "de"}])
    assert body["recipients"][0]["message_text"].startswith("Guten Tag")
    r = client.delete(f"/api/message-templates/{t2['id']}")
    assert r.status_code == 409 and r.json()["error"]["code"] == "in_use"
    assert client.patch(f"/api/message-templates/{t2['id']}", json={"status": "archived"}).status_code == 200
    assert client.delete(f"/api/message-templates/{t1['id']}").status_code == 204


# ---- mark sent and RFQ status ------------------------------------------------------------------

def test_mark_sent_creates_one_outgoing_interaction_with_followup(client, db):
    case, items = make_case(client, [{"name": "Valve"}, {"name": "Pump"}])
    party = supplier(client, country="AE")
    body = new_rfq(client, case, items[:1], [{"party_id": party["id"]}])
    recipient = body["recipients"][0]
    assert body["status"] == "draft"
    r = client.post(f"/api/rfqs/recipients/{recipient['id']}/mark-sent", json={"message_text": "Edited text"})
    assert r.status_code == 200, r.text
    sent = r.json()
    assert sent["status"] == "sent" and sent["sent_at"] and sent["message_text"] == "Edited text"
    rows = interactions_of(client, party["id"])
    assert len(rows) == 1
    row = client.get(f"/api/interactions/{rows[0]['id']}").json()
    assert (row["direction"], row["channel"], row["body"]) == ("out", "email", "Edited text")
    assert row["followups"][0]["action"] == "Chase RFQ response" and row["followups"][0]["due_on"] == soon(5)
    detail = client.get(f"/api/cases/{case['id']}").json()
    statuses = {i["name"]: i["status"] for i in detail["items"]}
    assert statuses == {"Valve": "rfq_sent", "Pump": "uncovered"}
    assert client.get(f"/api/rfqs/{body['id']}").json()["status"] == "sent"
    r = client.post(f"/api/rfqs/recipients/{recipient['id']}/mark-sent", json={})
    assert r.status_code == 409  # only once
    assert len(interactions_of(client, party["id"])) == 1


def test_mark_sent_with_owner_followup_or_without_respond_by(client):
    case, items = make_case(client, [{"name": "Bolt"}])
    party = supplier(client, country="IR")
    body = new_rfq(client, case, items, [{"party_id": party["id"]}], respond_by=None)
    rid = body["recipients"][0]["id"]
    r = client.post(f"/api/rfqs/recipients/{rid}/mark-sent", json={})
    assert r.status_code == 422 and r.json()["error"]["code"] == "next_action_required"
    r = client.post(f"/api/rfqs/recipients/{rid}/mark-sent",
                    json={"followup": {"action": "Call Reza", "due_on": soon(2)}})
    assert r.status_code == 200
    row = interactions_of(client, party["id"])[0]
    assert row["channel"] == "whatsapp"
    detail = client.get(f"/api/interactions/{row['id']}").json()
    assert detail["followups"][0]["action"] == "Call Reza"


def test_rfq_status_transitions(client, db):
    case, items = make_case(client, [{"name": "Cable"}])
    a, b, c = supplier(client), supplier(client), supplier(client)
    body = new_rfq(client, case, items, [{"party_id": p["id"]} for p in (a, b, c)])
    ra, rb, rc = (r["id"] for r in body["recipients"])
    assert body["status"] == "draft"
    client.post(f"/api/rfqs/recipients/{ra}/mark-sent", json={})
    client.post(f"/api/rfqs/recipients/{rb}/mark-sent", json={})
    assert client.get(f"/api/rfqs/{body['id']}").json()["status"] == "sent"
    client.post(f"/api/rfqs/recipients/{rb}/decline", json={"note": "no stock"})
    assert client.get(f"/api/rfqs/{body['id']}").json()["status"] == "sent"  # a is still waiting
    client.post("/api/offers", json=offer_body(case, a, [{"line_item_id": items[0]["id"], "unit_price": "3.2"}]))
    detail = client.get(f"/api/rfqs/{body['id']}").json()
    states = {r["party_id"]: r["status"] for r in detail["recipients"]}
    assert states == {a["id"]: "replied", b["id"]: "declined", c["id"]: "draft"}
    assert detail["status"] == "answered"
    assert client.delete(f"/api/rfqs/{body['id']}").json()["error"]["details"] == {"sent": 2, "offers": 1}
    r = client.post(f"/api/rfqs/{body['id']}/cancel", json={"reason": "case closed"})
    assert r.json()["status"] == "cancelled"
    draft = new_rfq(client, case, items, [{"party_id": c["id"]}])
    assert client.delete(f"/api/rfqs/{draft['id']}").status_code == 204
    assert audit_ops(db, "ops.rfq", draft["id"]) == ["I", "D"]


def test_expiry_handlers_change_only_status_and_record_events(client, db):
    case, items = make_case(client, [{"name": "Zinc"}])
    late, early = supplier(client), supplier(client)
    body = new_rfq(client, case, items, [{"party_id": late["id"]}, {"party_id": early["id"]}],
                   respond_by=(TODAY - timedelta(days=1)).isoformat())
    for r in body["recipients"]:
        client.post(f"/api/rfqs/recipients/{r['id']}/mark-sent",
                    json={"followup": {"action": "chase", "due_on": soon(1)}})
    off = client.post("/api/offers", json=offer_body(case, early, [{"line_item_id": items[0]["id"], "unit_price": "2"}],
                                                     valid_until=(TODAY - timedelta(days=2)).isoformat())).json()
    interactions_before = db.execute(text("SELECT count(*) FROM core.interaction")).scalar()
    from app.modules.sourcing.handlers import expire

    result = expire(db, {})
    db.commit()
    assert result["no_response"] >= 1 and result["offers_expired"] >= 1
    detail = client.get(f"/api/rfqs/{body['id']}").json()
    states = {r["party_id"]: r["status"] for r in detail["recipients"]}
    assert states == {late["id"]: "no_response", early["id"]: "replied"}
    assert detail["status"] == "answered"
    assert client.get(f"/api/offers/{off['id']}").json()["status"] == "expired"
    assert db.execute(text("SELECT count(*) FROM core.interaction")).scalar() == interactions_before
    assert events(client, "offer", off["id"], "status")[0]["data"]["new_status"] == "expired"
    assert any(e["data"].get("new_status") == "no_response" for e in events(client, "rfq", body["id"], "recipient_status"))
    # expired offers do not count as valid
    assert client.get(f"/api/cases/{case['id']}/coverage").json()["items"][0]["valid_offers"] == 0


def test_rfq_without_any_reply_becomes_no_response(client, db):
    case, items = make_case(client, [{"name": "Tin"}])
    party = supplier(client)
    body = new_rfq(client, case, items, [{"party_id": party["id"]}], respond_by=(TODAY - timedelta(days=1)).isoformat())
    client.post(f"/api/rfqs/recipients/{body['recipients'][0]['id']}/mark-sent",
                json={"followup": {"action": "chase", "due_on": soon(1)}})
    rfq.expire_no_reply(db)
    db.commit()
    assert client.get(f"/api/rfqs/{body['id']}").json()["status"] == "no_response"


def test_rfq_validation_and_duplicate_recipient_persists_nothing(client, db):
    case, items = make_case(client, [{"name": "Gasket"}])
    other_case, other_items = make_case(client, [{"name": "Other"}])
    party = supplier(client)
    r = client.post("/api/rfqs", json={"case_id": case["id"], "items": [{"line_item_id": other_items[0]["id"]}]})
    assert r.status_code == 422 and r.json()["error"]["code"] == "item_not_of_case"
    before = db.scalar(select(func.count()).select_from(Rfq))
    r = client.post("/api/rfqs", json={"case_id": case["id"], "items": [{"line_item_id": items[0]["id"]}],
                                       "recipients": [{"party_id": party["id"]}, {"party_id": party["id"]}]})
    assert r.status_code == 409 and r.json()["error"]["code"] == "duplicate_recipient"
    db.expire_all()
    assert db.scalar(select(func.count()).select_from(Rfq)) == before
    dnc = supplier(client, do_not_contact=True, do_not_contact_reason="asked")
    r = client.post("/api/rfqs", json={"case_id": case["id"], "items": [{"line_item_id": items[0]["id"]}],
                                       "recipients": [{"party_id": dnc["id"]}]})
    assert r.status_code == 422 and r.json()["error"]["code"] == "do_not_contact"


def test_xlsx_export(client):
    from io import BytesIO

    from openpyxl import load_workbook

    case, items = make_case(client, [{"name": "Copper wire", "quantity": "1500", "unit_code": "kg", "hs_code": "7408",
                                      "specs": [{"key": "Diameter", "value": "2.5", "unit": "mm"}]}])
    body = new_rfq(client, case, items, [])
    r = client.get(f"/api/rfqs/{body['id']}/export.xlsx")
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    ws = load_workbook(BytesIO(r.content)).active
    assert ws["A3"].value == "No" and ws["L3"].value == "Unit price"
    assert [ws.cell(4, c).value for c in (1, 2, 4, 5, 6, 11)] == [1, "Copper wire", 1500, "Kilogram",
                                                                  "Diameter: 2.5 mm", "7408"]
    fa = load_workbook(BytesIO(client.get(f"/api/rfqs/{body['id']}/export.xlsx", params={"language": "fa"}).content)).active
    assert fa["B3"].value == "کالا" and fa.sheet_view.rightToLeft


# ---- offers ---------------------------------------------------------------------------------------

def test_offer_creates_incoming_interaction_and_moves_items(client):
    case, items = make_case(client, [{"name": "Resin"}, {"name": "Hardener"}])
    party = supplier(client, country="CN")
    body = new_rfq(client, case, items, [{"party_id": party["id"]}])
    client.post(f"/api/rfqs/recipients/{body['recipients'][0]['id']}/mark-sent", json={})
    r = client.post("/api/offers", json=offer_body(case, party, [
        {"line_item_id": items[0]["id"], "unit_price": "1250", "unit_code": "t", "moq": "5"},
        {"line_item_id": items[1]["id"], "available": False}], channel="wechat", raw_text="报价 / quote",
        no_followup_reason=None, followup={"action": "Check docs", "due_on": soon(3)}))
    assert r.status_code == 201, r.text
    off = r.json()
    assert off["rfq_recipient_id"] == body["recipients"][0]["id"]  # matched automatically
    rows = [i for i in interactions_of(client, party["id"]) if i["direction"] == "in"]
    assert len(rows) == 1 and rows[0]["channel"] == "other" and rows[0]["subject"] == "WeChat"
    assert rows[0]["body"] == "报价 / quote"
    statuses = {i["name"]: i["status"] for i in client.get(f"/api/cases/{case['id']}").json()["items"]}
    assert statuses == {"Resin": "has_offer", "Hardener": "rfq_sent"}
    detail = client.get(f"/api/rfqs/{body['id']}").json()
    assert detail["recipients"][0]["status"] == "replied" and detail["status"] == "answered"


def test_offer_requires_next_action_and_valid_rows(client, db):
    case, items = make_case(client, [{"name": "Nut"}])
    party = supplier(client)
    body = offer_body(case, party, [{"line_item_id": items[0]["id"], "unit_price": "1"}], no_followup_reason=None)
    count = db.execute(text("SELECT count(*) FROM ops.supplier_offer")).scalar()
    r = client.post("/api/offers", json=body)
    assert r.status_code == 422 and r.json()["error"]["code"] == "next_action_required"
    r = client.post("/api/offers", json=offer_body(case, party, [{"line_item_id": items[0]["id"]}]))
    assert r.status_code == 422 and r.json()["error"]["code"] == "price_required"
    r = client.post("/api/offers", json=offer_body(case, party, [{"line_item_id": items[0]["id"], "unit_price": "1"}],
                                                   currency="XXQ"))
    assert r.status_code == 422
    db.expire_all()
    assert db.execute(text("SELECT count(*) FROM ops.supplier_offer")).scalar() == count
    assert interactions_of(client, party["id"]) == []


def test_superseding_keeps_history(client):
    case, items = make_case(client, [{"name": "Steel"}])
    party = supplier(client)
    first = client.post("/api/offers", json=offer_body(case, party, [{"line_item_id": items[0]["id"], "unit_price": "700"}])).json()
    second = client.post("/api/offers", json=offer_body(case, party, [{"line_item_id": items[0]["id"], "unit_price": "650"}],
                                                        supersedes_offer_id=first["id"])).json()
    old = client.get(f"/api/offers/{first['id']}").json()
    assert old["status"] == "superseded" and old["superseded_by_id"] == second["id"] and old["items"][0]["unit_price"] == "700.0000"
    assert len(client.get("/api/offers", params={"case_id": case["id"]}).json()) == 2
    assert events(client, "offer", first["id"], "superseded")[0]["data"]["superseded_by"] == second["id"]
    comp = client.get(f"/api/cases/{case['id']}/comparison").json()["items"][0]
    rows = [r for g in comp["groups"] for u in g["unit_groups"] for r in u["rows"]]
    assert [r["offer_id"] for r in rows] == [second["id"]]  # superseded offers are not valid
    other_party = supplier(client)
    r = client.post("/api/offers", json=offer_body(case, other_party, [{"line_item_id": items[0]["id"], "unit_price": "1"}],
                                                   supersedes_offer_id=second["id"]))
    assert r.status_code == 422 and r.json()["error"]["code"] == "supersede_mismatch"


# ---- comparison, coverage, selection ------------------------------------------------------------------

def test_group_offers_pure_helper():
    rows = [
        {"currency": "USD", "unit_group": "kind:weight", "incoterm": "FOB", "normalized_price": Decimal("1.2")},
        {"currency": "USD", "unit_group": "kind:weight", "incoterm": "FOB", "normalized_price": Decimal("1.1")},
        {"currency": "EUR", "unit_group": "kind:weight", "incoterm": "CIF", "normalized_price": Decimal("1.0")},
        {"currency": "EUR", "unit_group": "kind:weight", "incoterm": "FOB", "normalized_price": Decimal("0.9")},
        {"currency": "CNY", "unit_group": "kind:weight", "incoterm": "FOB", "normalized_price": Decimal("8")},
        {"currency": "CNY", "unit_group": "unit:drum", "incoterm": "FOB", "normalized_price": Decimal("900")},
    ]
    groups = {g["currency"]: g for g in coverage.group_offers(rows)}
    assert list(groups) == ["CNY", "EUR", "USD"]
    usd = groups["USD"]["unit_groups"][0]["rows"]
    assert [r["lowest"] for r in usd] == [False, True] and all(r["badges"] == [] for r in usd)
    eur = groups["EUR"]["unit_groups"][0]
    assert eur["comparable"] is False and all(r["badges"] == ["incoterm_differs"] and not r["lowest"] for r in eur["rows"])
    cny = [r for u in groups["CNY"]["unit_groups"] for r in u["rows"]]
    assert all("unit_differs" in r["badges"] for r in cny) and all(r["lowest"] for r in cny)  # each alone in its group


def test_comparison_and_selection_via_api(client):
    case, items = make_case(client, [{"name": "PVC resin", "quantity": "40", "unit_code": "t"}])
    item = items[0]
    a, b, c, d = (supplier(client, name=f"S{n} {tag()}") for n in range(4))
    oa = client.post("/api/offers", json=offer_body(case, a, [{"line_item_id": item["id"], "unit_price": "980", "unit_code": "t"}],
                                                    incoterm="CFR")).json()
    ob = client.post("/api/offers", json=offer_body(case, b, [{"line_item_id": item["id"], "unit_price": "0.95", "unit_code": "kg"}],
                                                    incoterm="CFR", valid_until=soon(3))).json()
    oc = client.post("/api/offers", json=offer_body(case, c, [{"line_item_id": item["id"], "unit_price": "820", "unit_code": "t"}],
                                                    currency="EUR", incoterm="FOB")).json()
    od = client.post("/api/offers", json=offer_body(case, d, [{"line_item_id": item["id"], "unit_price": "40", "unit_code": "bag"}],
                                                    currency="EUR", incoterm="FOB")).json()
    comp = client.get(f"/api/cases/{case['id']}/comparison", params={"item_id": item["id"]}).json()["items"][0]
    groups = {g["currency"]: g for g in comp["groups"]}
    usd = groups["USD"]["unit_groups"]
    assert len(usd) == 1  # t and kg are convertible (same kind)
    by_offer = {r["offer_id"]: r for r in usd[0]["rows"]}
    assert by_offer[ob["id"]]["lowest"] and not by_offer[oa["id"]]["lowest"]  # 950/t < 980/t
    assert by_offer[ob["id"]]["expiring"] and by_offer[ob["id"]]["days_left"] == 3
    eur_rows = {r["offer_id"]: r for u in groups["EUR"]["unit_groups"] for r in u["rows"]}
    assert "unit_differs" in eur_rows[oc["id"]]["badges"] and "unit_differs" in eur_rows[od["id"]]["badges"]
    cov = client.get(f"/api/cases/{case['id']}/coverage").json()
    assert cov["items"][0]["valid_offers"] == 4 and cov["gate"]["ok"] is False
    best = {(b["currency"], b["unit_group"]): b for b in cov["items"][0]["best_prices"]}
    assert best[("USD", "kind:weight")]["offer_id"] == ob["id"]
    # select, then change with a reason
    ia = next(i["id"] for i in client.get(f"/api/offers/{oa['id']}").json()["items"])
    ib = next(i["id"] for i in client.get(f"/api/offers/{ob['id']}").json()["items"])
    r = client.post(f"/api/line-items/{item['id']}/select", json={"offer_item_id": ib})
    assert r.json()["status"] == "selected" and r.json()["selected_offer_item_id"] == ib
    assert client.get(f"/api/cases/{case['id']}/coverage").json()["gate"]["ok"] is True
    r = client.post(f"/api/line-items/{item['id']}/select", json={"offer_item_id": ia})
    assert r.status_code == 422 and r.json()["error"]["code"] == "reason_required"
    r = client.post(f"/api/line-items/{item['id']}/select", json={"offer_item_id": ia, "reason": "better lead time"})
    assert r.json()["selected_offer_item_id"] == ia
    ev = events(client, "line_item", item["id"], "select_offer")[0]["data"]
    assert ev == {"old_offer_item_id": ib, "new_offer_item_id": ia, "reason": "better lead time",
                  "old_status": "selected", "new_status": "selected"}
    sel = client.get(f"/api/cases/{case['id']}/comparison").json()["items"][0]
    assert [r["offer_item_id"] for g in sel["groups"] for u in g["unit_groups"] for r in u["rows"] if r["selected"]] == [ia]
    r = client.delete(f"/api/offers/{oa['id']}")
    assert r.status_code == 409 and r.json()["error"]["details"]["selected_items"] == 1
    assert client.post(f"/api/offers/{oc['id']}/status", json={"status": "withdrawn"}).json()["status"] == "withdrawn"
    assert client.delete(f"/api/offers/{od['id']}").status_code == 204


def test_gate_with_no_coverage_and_cancelled_items(client):
    case, items = make_case(client, [{"name": "A"}, {"name": "B"}, {"name": "C"}])
    party = supplier(client)
    off = client.post("/api/offers", json=offer_body(case, party, [{"line_item_id": items[0]["id"], "unit_price": "5"}])).json()
    gate = client.get(f"/api/cases/{case['id']}/coverage").json()["gate"]
    assert {m["name"] for m in gate["missing"]} == {"A", "B", "C"}
    client.post(f"/api/line-items/{items[0]['id']}/select", json={"offer_item_id": off["items"][0]["id"]})
    client.post(f"/api/line-items/{items[1]['id']}/no-coverage", json={"reason": "not produced anymore"})
    assert {m["name"] for m in client.get(f"/api/cases/{case['id']}/coverage").json()["gate"]["missing"]} == {"C"}
    client.post(f"/api/line-items/{items[2]['id']}/cancel", json={"reason": "dropped"})
    assert client.get(f"/api/cases/{case['id']}/coverage").json()["gate"] == {"ok": True, "missing": []}


# ---- capabilities ----------------------------------------------------------------------------------

def test_reference_seed_counts_hierarchy_and_fa_fallback(db):
    counts = dict(db.execute(text("SELECT level::text, count(*) FROM ref.hs_code GROUP BY 1")).all())
    assert counts == {"2": 96, "4": 1228}
    assert db.execute(text("SELECT count(*) FROM ref.hs_code h WHERE level = 4 AND NOT EXISTS "
                           "(SELECT 1 FROM ref.hs_code p WHERE p.code = h.parent_code AND p.level = 2)")).scalar() == 0
    counts = dict(db.execute(text("SELECT level, count(*) FROM ref.isic_code GROUP BY 1")).all())
    assert counts == {"section": 21, "division": 88}
    assert db.execute(text("SELECT parent_code FROM ref.isic_code WHERE code = '20'")).scalar() == "C"
    assert db.execute(text("SELECT edition FROM ref.hs_code WHERE code = '3907'")).scalar() == "HS 2022"


def test_code_search_tree_and_missing_persian_title(client):
    rows = client.get("/api/capabilities/codes", params={"kind": "hs", "q": "3907"}).json()
    row = next(r for r in rows if r["code"] == "3907")
    assert row["parent"]["code"] == "39" and row["parent"]["title_fa"] == "پلاستیک و اشیای ساخته‌شده از آن"
    assert row["title_fa"] is None and row["title_fa_missing"] is True and row["title_en"].startswith("Polyacetals")
    assert any(r["code"] == "39" for r in client.get("/api/capabilities/codes", params={"kind": "hs", "q": "پلاستيک"}).json())
    top = client.get("/api/capabilities/codes", params={"kind": "isic"}).json()
    assert [r["code"] for r in top][:3] == ["A", "B", "C"]
    children = client.get("/api/capabilities/codes", params={"kind": "isic", "parent": "C"}).json()
    assert len(children) == 24 and children[10]["code"] == "20"


def test_capabilities_unique_multiselect_and_tag_normalization(client, db):
    party = supplier(client)
    pid = party["id"]
    r = client.post(f"/api/capabilities/party/{pid}", json={"kind": "hs", "code": "39.07"})
    assert r.status_code == 201 and r.json()[0]["code"] == "3907"
    r = client.post(f"/api/capabilities/party/{pid}", json={"kind": "hs", "code": "3907"})
    assert r.status_code == 409 and r.json()["error"]["code"] == "duplicate_capability"
    assert client.post(f"/api/capabilities/party/{pid}", json={"kind": "hs", "code": "9999"}).json()["error"]["code"] == "unknown_code"
    client.post(f"/api/capabilities/party/{pid}", json={"kind": "tag", "tag_text": "  PE100   Grade "})
    r = client.post(f"/api/capabilities/party/{pid}", json={"kind": "tag", "tag_text": "pe100 grade"})
    assert r.status_code == 409
    r = client.put(f"/api/capabilities/party/{pid}", json={"items": [
        {"kind": "hs", "code": "3907"}, {"kind": "isic", "code": "C20"}, {"kind": "tag", "tag_text": "كيميا"}]})
    assert r.status_code == 200
    caps = {(c["kind"], c["code"] or c["tag_text"]) for c in r.json()}
    assert caps == {("hs", "3907"), ("isic", "20"), ("tag", "كيميا")}
    tag_row = db.execute(text("SELECT tag_norm FROM core.party_capability WHERE party_id = :p AND kind = 'tag'"),
                         {"p": pid}).scalar()
    assert tag_row == "کیمیا"  # Arabic yeh/kaf folded
    again = client.put(f"/api/capabilities/party/{pid}", json={"items": [
        {"kind": "hs", "code": "3907"}, {"kind": "isic", "code": "20"}, {"kind": "tag", "tag_text": "كيميا"}]}).json()
    assert {c["id"] for c in again} == {c["id"] for c in r.json()}  # round-trip keeps rows
    adds = events(client, "party", pid, "capability_add")
    assert {(e["data"]["kind"], e["data"]["code"] or e["data"]["tag"]) for e in adds} >= {("isic", "20")}
    assert audit_ops(db, "core.party_capability", r.json()[0]["id"])[0] == "I"


def test_party_filter_by_hs_isic_tag(client):
    heading, chapter, other, isic, tagged = (supplier(client) for _ in range(5))
    client.post(f"/api/capabilities/party/{heading['id']}", json={"kind": "hs", "code": "3907"})
    client.post(f"/api/capabilities/party/{chapter['id']}", json={"kind": "hs", "code": "39"})
    client.post(f"/api/capabilities/party/{other['id']}", json={"kind": "hs", "code": "7403"})
    client.post(f"/api/capabilities/party/{isic['id']}", json={"kind": "isic", "code": "C"})
    word = f"brand{tag()}"
    client.post(f"/api/capabilities/party/{tagged['id']}", json={"kind": "tag", "tag_text": word.upper()})

    def ids(spec):
        return {p["id"] for p in client.get("/api/parties", params={"capability": spec, "page_size": 200}).json()["items"]}

    assert {heading["id"], chapter["id"]} <= ids("hs:3907") and other["id"] not in ids("hs:3907")
    assert {heading["id"], chapter["id"]} <= ids("hs:39")
    assert isic["id"] in ids("isic:C20") and isic["id"] in ids("isic:C")
    assert ids(f"tag:{word}") == {tagged["id"]}
    assert client.get("/api/parties", params={"capability": "xx:1"}).status_code == 422


def test_rfq_supplier_suggestions_heading_then_chapter(client):
    case, items = make_case(client, [{"name": "POM", "hs_code": "390710"}])
    heading, chapter, unrelated = supplier(client), supplier(client), supplier(client)
    client.post(f"/api/capabilities/party/{heading['id']}", json={"kind": "hs", "code": "3907"})
    client.post(f"/api/capabilities/party/{chapter['id']}", json={"kind": "hs", "code": "3901"})
    client.post(f"/api/capabilities/party/{unrelated['id']}", json={"kind": "hs", "code": "7403"})
    rows = client.get("/api/rfqs/supplier-suggestions", params={"item_ids": [items[0]["id"]]}).json()
    by_party = {r["party_id"]: r for r in rows}
    assert by_party[heading["id"]]["match"] == "heading" and by_party[chapter["id"]]["match"] == "chapter"
    assert unrelated["id"] not in by_party
    order = [r["party_id"] for r in rows]
    assert order.index(heading["id"]) < order.index(chapter["id"])
    assert by_party[heading["id"]]["supplier_stage"] == "identified"
    rows = client.get("/api/rfqs/supplier-suggestions",
                      params={"item_ids": [items[0]["id"]], "exclude": [heading["id"]]}).json()
    assert heading["id"] not in {r["party_id"] for r in rows}
    body = new_rfq(client, case, items, [])
    assert body["recipients"] == []  # suggestions never add anyone


def test_party_search_finds_capability(client):
    party = supplier(client, name=f"Nameless {tag()}")
    word = f"cap{tag()}"
    client.post(f"/api/capabilities/party/{party['id']}", json={"kind": "tag", "tag_text": word})
    client.post(f"/api/capabilities/party/{party['id']}", json={"kind": "hs", "code": "2815"})
    for q in (word, "2815"):
        groups = client.get("/api/search", params={"q": q, "types": "party", "limit": 50}).json()["groups"]
        assert party["id"] in [h["id"] for g in groups for h in g["items"]], q


def test_edition_replaceable_without_breaking_links(client):
    """Replacing the reference rows (new edition, changed title, removed code) never breaks party links."""
    party = supplier(client)
    client.post(f"/api/capabilities/party/{party['id']}", json={"kind": "hs", "code": "8524"})
    conninfo, db_name = superuser_conninfo(), get_config().db_name
    with dbadmin.superuser_conn(conninfo, db_name) as conn:
        original = conn.execute("SELECT code, level, parent_code, title_en, title_fa, edition FROM ref.hs_code "
                                "WHERE code = '8524'").fetchone()
        conn.execute("UPDATE ref.hs_code SET edition = 'HS 2028 (test)', title_en = 'Flat panel modules (new)' "
                     "WHERE code = '8524'")
        conn.commit()
    try:
        caps = client.get(f"/api/capabilities/party/{party['id']}").json()
        assert caps[0]["info"]["edition"] == "HS 2028 (test)" and caps[0]["info"]["title_en"] == "Flat panel modules (new)"
        with dbadmin.superuser_conn(conninfo, db_name) as conn:
            conn.execute("DELETE FROM ref.hs_code WHERE code = '8524'")
            conn.commit()
        caps = client.get(f"/api/capabilities/party/{party['id']}").json()
        assert caps[0]["code"] == "8524" and caps[0]["info"] == {"code": "8524", "unknown": True}
        assert party["id"] in {p["id"] for p in client.get("/api/parties", params={"capability": "hs:8524"}).json()["items"]}
    finally:
        with dbadmin.superuser_conn(conninfo, db_name) as conn:
            conn.execute("DELETE FROM ref.hs_code WHERE code = '8524'")
            conn.execute("INSERT INTO ref.hs_code (code, level, parent_code, title_en, title_fa, edition) "
                         "VALUES (%s, %s, %s, %s, %s, %s)", tuple(original))
            conn.commit()
    caps = client.get(f"/api/capabilities/party/{party['id']}").json()
    assert caps[0]["info"]["edition"] == "HS 2022"


def test_unknown_code_is_reported_not_broken(db):
    from app.modules.sourcing import capabilities

    assert capabilities.code_info(db, "hs", "0000") == {"code": "0000", "unknown": True}


# ---- catalog ----------------------------------------------------------------------------------------

def test_catalog_duplicate_warning_and_alias_search(client):
    name = f"Polypropylene homopolymer {tag()}"
    first = client.post("/api/products", json={"name_en": name, "hs_code": "3902", "aliases": [f"PP-H {tag()}", "يكي"]})
    assert first.status_code == 201
    pid = first.json()["product"]["id"]
    r = client.post("/api/products", json={"name_en": name.replace("homopolymer", "homopolymr")})
    assert r.status_code == 200 and r.json()["created"] is False
    assert pid in [d["id"] for d in r.json()["possible_duplicates"]]
    r = client.post("/api/products", json={"name_en": name.replace("homopolymer", "homopolymr"), "confirm_duplicate": True})
    assert r.status_code == 201
    hits = client.get("/api/products", params={"q": "یکی"}).json()["items"]  # Persian yeh/kaf finds Arabic alias
    assert pid in [p["id"] for p in hits]
    groups = client.get("/api/search", params={"q": "يكي", "types": "product"}).json()["groups"]
    assert pid in [h["id"] for g in groups for h in g["items"]]
    assert pid in [p["id"] for p in client.get("/api/products", params={"q": "3902"}).json()["items"]]


def test_product_from_item_and_delete_rules(client, db):
    case, items = make_case(client, [{"name": f"اسید سولفوریک {tag()}", "unit_code": "kg", "hs_code": "2807"}])
    r = client.post(f"/api/line-items/{items[0]['id']}/create-product", json={})
    assert r.status_code == 201, r.text
    pid = r.json()["product_id"]
    product = client.get(f"/api/products/{pid}").json()
    assert product["name_fa"] == items[0]["name"] and product["name_en"] is None and product["hs_code"] == "2807"
    assert r.json()["item"]["product_id"] == pid
    r = client.delete(f"/api/products/{pid}")
    assert r.status_code == 409 and r.json()["error"]["details"] == {"line_items": 1}
    assert client.post(f"/api/products/{pid}/archive").json()["status"] == "archived"
    loose = client.post("/api/products", json={"name_en": f"Loose {tag()}", "aliases": ["x1"]}).json()["product"]
    assert client.delete(f"/api/products/{loose['id']}").status_code == 204
    assert audit_ops(db, "core.product", loose["id"]) == ["I", "D"]


def test_product_price_memory_and_demand_on_fixed_data(client):
    product = client.post("/api/products", json={"name_en": f"Test glycol {tag()}", "confirm_duplicate": True}).json()["product"]
    buyer = make_party(client, roles=["buyer"])
    prices = [("USD", "100", "t"), ("USD", "0.12", "kg"), ("USD", "90", "t"), ("EUR", "80", "t"), ("USD", "500", "drum")]
    for currency, price, unit in prices:
        case, items = make_case(client, [{"name": "glycol", "product_id": product["id"], "quantity": "2", "unit_code": "t"}],
                                party_id=buyer["id"])
        party = supplier(client)
        client.post("/api/offers", json=offer_body(case, party, [{"line_item_id": items[0]["id"], "unit_price": price,
                                                                  "unit_code": unit}], currency=currency))
    page = client.get(f"/api/products/{product['id']}").json()
    summary = {(s["currency"], s["unit_group"]): s for s in page["prices"]["summary"]}
    usd = summary[("USD", "kind:weight")]["all_time"]
    # per kg: 0.10, 0.12, 0.09
    assert usd["count"] == 3 and Decimal(usd["min"]) == Decimal("0.09") and Decimal(usd["median"]) == Decimal("0.1")
    assert Decimal(usd["max"]) == Decimal("0.12") and summary[("USD", "kind:weight")]["last_12_months"]["count"] == 3
    assert summary[("EUR", "kind:weight")]["all_time"]["count"] == 1
    assert summary[("USD", "unit:drum")]["all_time"]["count"] == 1
    assert len(page["prices"]["offers"]) == 5
    demand = page["demand"]
    assert demand["cases"] == 5 and demand["inquiries"] == 0
    month = datetime.now(UTC).astimezone().strftime("%Y-%m")
    assert {(m["month"], m["unit_group"]): Decimal(m["quantity"]) for m in demand["quantity_by_month"]} == \
        {(month, "kind:weight"): Decimal("10000")}  # 5 x 2 t in kg
    assert demand["buyers"][0]["party_id"] == buyer["id"] and demand["buyers"][0]["requests"] == 5
    assert len(page["suppliers_by_offers"]) == 5
    listed = client.get("/api/products", params={"q": product["name_en"], "sort": "offers"}).json()["items"]
    assert listed[0]["offers"] == 5 and listed[0]["last_requested"]


# ---- command center (isolated database: exact numbers) -------------------------------------------------

def test_command_center_sourcing_numbers_on_fixed_dataset():
    from alembic import command

    from app.modules.parties import service as parties
    from app.modules.sourcing import inquiries as inquiries_service
    from app.modules.sourcing import items as items_service
    from app.modules.system.seed import seed_base_data

    name = "boms_test_cc_sourcing"
    create_empty_database(name)
    url = get_config().db_url(name)
    command.upgrade(alembic_config(url), "head")
    engine = create_engine(url)
    today = clock.today()
    try:
        with Session(engine) as db:
            seed_base_data(db)
            buyer, _ = parties.create_party(db, {"name_en": "Buyer", "roles": ["buyer"], "confirm_duplicate": True}, "script")
            sup, _ = parties.create_party(db, {"name_en": "Supplier", "roles": ["supplier"], "confirm_duplicate": True},
                                          "script")
            case = inquiries_service.create_case(db, {"title": "Open case", "party_id": buyer.id}, "script")
            closed = inquiries_service.create_case(db, {"title": "Closed case"}, "script")
            a, b, c, d = items_service.add_items(db, None, case.id, [{"name": n} for n in "ABCD"], "script")
            items_service.add_items(db, None, closed.id, [{"name": "X"}], "script")
            inquiries_service.update_case(db, closed.id, {"status": "closed"}, "script")
            body = rfq.create_rfq(db, {"case_id": case.id, "respond_by": today - timedelta(days=2),
                                       "items": [{"line_item_id": a.id}, {"line_item_id": b.id}],
                                       "recipients": [{"party_id": sup.id}]}, "script")
            rec = rfq.recipients_of(db, body.id)[0]
            rfq.mark_sent(db, rec.id, {"followup": {"action": "chase", "due_on": today}}, "script")
            offer = offers.record_offer(db, {"case_id": case.id, "party_id": sup.id, "channel": "email",
                                             "currency": "USD", "valid_until": today + timedelta(days=3),
                                             "items": [{"line_item_id": b.id, "unit_price": Decimal("5")}],
                                             "no_followup_reason": "fyi"}, "script")
            # the recipient replied; another supplier's request is overdue
            sup2, _ = parties.create_party(db, {"name_en": "Supplier 2", "roles": ["supplier"], "confirm_duplicate": True},
                                           "script")
            rfq.add_recipient(db, body.id, {"party_id": sup2.id}, "script")
            rec2 = [r for r in rfq.recipients_of(db, body.id) if r.party_id == sup2.id][0]
            rfq.mark_sent(db, rec2.id, {"followup": {"action": "chase", "due_on": today}}, "script")
            items_service.set_no_coverage(db, c.id, "not available", "script")
            db.commit()
            result = dashboard.summary(db, today)
            assert [i["name"] for i in result["items_without_rfq"]] == ["D"]  # C is marked no coverage
            assert [r["party_name_en"] for r in result["rfq_overdue"]] == ["Supplier 2"]
            assert [o["offer_id"] for o in result["offers_expiring"]] == [str(offer.id)]
            assert result["offers_expiring"][0]["days_left"] == 3
            assert result["cases_not_covered"] == [{"case_id": str(case.id), "case_title": "Open case", "items": 4,
                                                    "not_covered": 3}]
            # A rfq_sent, B has_offer, C no coverage, D uncovered -> 3 of 4
            assert (result["open_items"], result["open_items_with_status"], result["coverage_status_percent"]) == (4, 3, 75.0)
            from app.modules.parties import command_center

            assert command_center.summary(db)["sourcing"]["coverage_status_percent"] == 75.0
    finally:
        engine.dispose()
        dbadmin.drop_database(superuser_conninfo(), name)


def test_command_center_api_has_sourcing_section(client):
    body = client.get("/api/command-center").json()
    assert {"items_without_rfq", "rfq_overdue", "offers_expiring", "cases_not_covered",
            "coverage_status_percent"} <= set(body["sourcing"])


# ---- search ------------------------------------------------------------------------------------------

def test_search_providers_with_arabic_yeh_kaf(client):
    word = f"كيك{tag()}"  # Arabic kaf/yeh
    persian = word.replace("ك", "ک").replace("ي", "ی")
    inq = client.post("/api/inquiries", json={"source_channel": "email", "subject": f"استعلام {persian}",
                                              "raw_text": f"متن {persian}", "items": [{"name": f"قلم {persian}"}]}).json()
    case = client.post("/api/cases", json={"title": f"پرونده {persian}"}).json()
    citem = client.post("/api/line-items", json={"case_id": case["id"], "name": "Item X"}).json()
    party = make_party(client, name_fa=f"تامین {persian}", roles=["supplier"])
    body = new_rfq(client, case, [citem], [{"party_id": party["id"]}], title=f"RFQ {persian}")
    off = client.post("/api/offers", json=offer_body(case, party, [{"line_item_id": citem["id"], "unit_price": "1"}],
                                                     raw_text=f"پیشنهاد {persian}")).json()
    groups = {g["type"]: [h["id"] for h in g["items"]]
              for g in client.get("/api/search", params={"q": word, "limit": 20}).json()["groups"]}
    assert inq["id"] in groups["inquiry"] and case["id"] in groups["case"]
    assert inq["items"][0]["id"] in groups["line_item"] and body["id"] in groups["rfq"]
    assert off["id"] in groups["offer"]


# ---- hardening, audit, backup ----------------------------------------------------------------------------

@pytest.mark.parametrize("method,path,body", [
    ("post", "/api/rfqs", {"case_id": "00000000-0000-0000-0000-000000000000", "items": [], "bogus": 1}),
    ("post", "/api/offers", {"bogus": 1}),
    ("post", "/api/message-templates", {"language": "en", "channel": "email", "name": "x", "body_template": "x", "bogus": 1}),
    ("post", "/api/products", {"name_en": "x", "bogus": 1}),
    ("post", "/api/products/categories", {"name_en": "x", "bogus": 1}),
    ("put", "/api/capabilities/party/00000000-0000-0000-0000-000000000000", {"items": [], "bogus": 1}),
    ("post", "/api/rfqs/recipients/00000000-0000-0000-0000-000000000000/mark-sent", {"bogus": 1}),
    ("patch", "/api/rfqs/recipients/00000000-0000-0000-0000-000000000000", {"bogus": 1}),
    ("post", "/api/line-items/00000000-0000-0000-0000-000000000000/select", {"offer_item_id": None, "bogus": 1}),
    ("put", "/api/inquiries/00000000-0000-0000-0000-000000000000/decision", {"light": "green", "reason": "x", "bogus": 1}),
    ("patch", "/api/offers/items/00000000-0000-0000-0000-000000000000", {"bogus": 1}),
])
def test_every_new_request_model_rejects_unknown_fields(client, method, path, body):
    r = getattr(client, method)(path, json=body)
    assert r.status_code == 422 and r.json()["error"]["code"] == "validation_error"


NEW_TABLES = ["ops.payment_term", "ops.item_status", "ops.inquiry", "ops.decision", "ops.case", "ops.line_item",
              "ops.rfq", "ops.rfq_item", "ops.rfq_recipient", "ops.message_template", "ops.supplier_offer",
              "ops.offer_item", "core.product_category", "core.product", "core.product_alias",
              "core.party_capability", "ref.hs_code", "ref.isic_code"]


def test_audit_trigger_on_every_new_table_and_rows_written(client, db):
    triggers = set(db.execute(text(
        "SELECT n.nspname || '.' || c.relname FROM pg_trigger t JOIN pg_class c ON c.oid = t.tgrelid "
        "JOIN pg_namespace n ON n.oid = c.relnamespace WHERE t.tgname LIKE '%\\_audit'")).scalars())
    assert set(NEW_TABLES) <= triggers
    # a full flow writes audit rows for the mutable tables
    case, items = make_case(client, [{"name": "Audit item"}])
    party = supplier(client)
    body = new_rfq(client, case, items, [{"party_id": party["id"]}])
    off = client.post("/api/offers", json=offer_body(case, party, [{"line_item_id": items[0]["id"], "unit_price": "1"}])).json()
    product = client.post("/api/products", json={"name_en": f"Audit product {tag()}", "aliases": ["aud"],
                                                 "confirm_duplicate": True}).json()["product"]
    category = client.post("/api/products/categories", json={"name_en": f"Cat {tag()}"}).json()
    cap = client.put(f"/api/capabilities/party/{party['id']}", json={"items": [{"kind": "hs", "code": "3907"}]}).json()[0]
    template = client.post("/api/message-templates", json={"language": "en", "channel": "letter", "name": f"L {tag()}",
                                                           "body_template": "Dear {contact_name}"}).json()
    inq = client.post("/api/inquiries", json={"source_channel": "email"}).json()
    client.put(f"/api/inquiries/{inq['id']}/decision", json={"light": "green", "reason": "ok"})
    decision_id = client.get(f"/api/inquiries/{inq['id']}").json()["decision"]["id"]
    rows = {
        "ops.case": case["id"], "ops.line_item": items[0]["id"], "ops.rfq": body["id"],
        "ops.rfq_item": body["items"][0]["id"], "ops.rfq_recipient": body["recipients"][0]["id"],
        "ops.supplier_offer": off["id"], "ops.offer_item": off["items"][0]["id"], "core.product": product["id"],
        "core.product_alias": product["aliases"][0]["id"], "core.product_category": category["id"],
        "core.party_capability": cap["id"], "ops.message_template": template["id"], "ops.inquiry": inq["id"],
        "ops.decision": decision_id,
    }
    for table, row_id in rows.items():
        assert audit_ops(db, table, row_id)[:1] == ["I"], table


def test_backup_and_restore_test_include_sourcing_tables(client):
    from app.modules.system import backup

    case, items = make_case(client, [{"name": "Backup item", "specs": [{"key": "k", "value": "v"}]}])
    party = supplier(client)
    new_rfq(client, case, items, [{"party_id": party["id"]}])
    client.post("/api/offers", json=offer_body(case, party, [{"line_item_id": items[0]["id"], "unit_price": "1"}]))
    doc = upload(client, "offer.txt", unique_bytes("offer doc")).json()
    client.post("/api/products", json={"name_en": f"Backup product {tag()}", "confirm_duplicate": True})
    run_id = backup.run_backup("manual", source="script")
    result = backup.restore_test(run_id, source="script")
    assert result["status"] == "passed", result
    # 18 tables up to v0.2.1 + 18 new ones (ops."case" included: the count query quotes reserved names)
    assert result["details"]["tables_checked"] == 36 and result["details"]["row_count_mismatches"] == {}
    assert doc["id"]
