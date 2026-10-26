"""RFQs: items + recipients, one rendered text per recipient, "mark as sent", xlsx export.

BOMS never sends anything. "Mark as sent" records the fact: the recipient becomes ``sent``, an OUTGOING interaction
(parties public service) is created with the exact text and a follow-up (default "Chase RFQ response" on
``respond_by``), and the RFQ's items' status is recomputed (``uncovered`` -> ``rfq_sent``).
"""

import io
import uuid
from datetime import date
from urllib.parse import quote

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.errors import AppError, Conflict, NotFound
from app.core.time import ensure_aware
from app.modules.events import service as events
from app.modules.parties import clock
from app.modules.parties import interactions as party_interactions
from app.modules.parties import service as parties
from app.modules.sourcing import items as items_service
from app.modules.sourcing import templates
from app.modules.sourcing.common import flush, jsonable, row_snapshot
from app.modules.sourcing.inquiries import get_case
from app.modules.sourcing.models import LineItem, Rfq, RfqItem, RfqRecipient, SupplierOffer
from app.modules.system import settings
from app.modules.system.models import Country, Unit

DOMESTIC_COUNTRY = "IR"
CHANNEL_TO_INTERACTION = {"email": "email", "whatsapp": "whatsapp", "letter": "other", "phone": "call",
                          "other": "other"}
DEFAULT_CHASE_ACTION = {"en": "Chase RFQ response", "fa": "پیگیری پاسخ استعلام"}
WAITING = "sent"


def get_rfq(db: Session, rfq_id: uuid.UUID) -> Rfq:
    row = db.get(Rfq, rfq_id)
    if row is None:
        raise NotFound("not_found", "RFQ not found")
    return row


def get_recipient(db: Session, recipient_id: uuid.UUID) -> RfqRecipient:
    row = db.get(RfqRecipient, recipient_id)
    if row is None:
        raise NotFound("not_found", "RFQ recipient not found")
    return row


def rfq_items(db: Session, rfq_id: uuid.UUID) -> list[tuple[RfqItem, LineItem]]:
    return list(db.execute(select(RfqItem, LineItem).join(LineItem, LineItem.id == RfqItem.line_item_id)
                           .where(RfqItem.rfq_id == rfq_id).order_by(LineItem.position, LineItem.created_at)).all())


def recipients_of(db: Session, rfq_id: uuid.UUID) -> list[RfqRecipient]:
    return list(db.scalars(select(RfqRecipient).where(RfqRecipient.rfq_id == rfq_id)
                           .order_by(RfqRecipient.created_at)))


def defaults_for_party(party) -> dict:
    """Domestic (Iran) suppliers: WhatsApp in Persian; foreign ones: e-mail in English. Editable afterwards."""
    domestic = (party.country_iso2 or "").upper() == DOMESTIC_COUNTRY
    return {"channel": "whatsapp" if domestic else "email", "language": "fa" if domestic else "en"}


# ---- rendering ------------------------------------------------------------------------------

def _name(lang: str, fa: str | None, en: str | None) -> str:
    return ((fa or en) if lang == "fa" else (en or fa)) or ""


def table_rows(db: Session, rfq_id: uuid.UUID, language: str) -> list[dict]:
    units = {u.code: u for u in db.scalars(select(Unit))}
    countries = {}
    rows = []
    for rfq_item, item in rfq_items(db, rfq_id):
        unit_code = rfq_item.unit_code or (None if rfq_item.unit_text else item.unit_code)
        unit = units.get(unit_code)
        unit_label = (_name(language, unit.name_fa, unit.name_en) if unit else None) or rfq_item.unit_text or item.unit_text
        destination = None
        if item.destination_country_iso2:
            iso = item.destination_country_iso2
            if iso not in countries:
                countries[iso] = db.get(Country, iso)
            c = countries[iso]
            destination = _name(language, c.name_fa, c.name_en) if c else iso
        incoterm = " ".join(x for x in (item.incoterm, item.incoterm_place) if x) or None
        rows.append({"name": item.name, "quantity": rfq_item.quantity or item.quantity, "unit": unit_label,
                     "specs": item.specs, "packaging": item.packaging, "destination": destination,
                     "incoterm": incoterm, "required_by": templates.format_date(item.required_by, language) or None,
                     "note": rfq_item.note})
    return rows


def signature(db: Session, language: str) -> str:
    sig = settings.get(db, "company.signature") or {}
    text = sig.get(language) or sig.get("en") or ""
    return text or settings.get(db, "company.name") or ""


def render_values(db: Session, rfq: Rfq, recipient: RfqRecipient) -> dict[str, str]:
    lang = recipient.language
    party = parties.get_party(db, recipient.party_id)
    supplier = _name(lang, party.name_fa, party.name_en)
    contact_name = supplier
    if recipient.contact_id:
        contact = parties.get_contact(db, recipient.contact_id)
        contact_name = _name(lang, contact.name_fa, contact.name_en) or supplier
    return {
        "supplier_name": supplier,
        "contact_name": contact_name,
        "case_title": get_case(db, rfq.case_id).title,
        "items_table": templates.items_table(table_rows(db, rfq.id, lang), lang),
        "respond_by": templates.format_date(rfq.respond_by, lang),
        "signature": signature(db, lang),
    }


def render_recipient(db: Session, rfq: Rfq, recipient: RfqRecipient) -> None:
    template = (templates.get_template(db, recipient.template_id) if recipient.template_id
                else templates.default_template(db, recipient.language, recipient.channel))
    if template is None:
        raise AppError("no_template", "No message template for this language and channel",
                       {"language": recipient.language, "channel": recipient.channel}, status=422)
    recipient.template_id = template.id
    values = render_values(db, rfq, recipient)
    recipient.subject = templates.render(template.subject_template, values).strip() or None
    recipient.message_text = templates.render(template.body_template, values)


def email_of(db: Session, recipient: RfqRecipient) -> str | None:
    contacts = parties.contacts_of(db, recipient.party_id)
    if recipient.contact_id:
        contacts = [c for c in contacts if c.id == recipient.contact_id] + [c for c in contacts
                                                                              if c.id != recipient.contact_id]
    methods = parties.methods_of(db, [c.id for c in contacts])
    for contact in contacts:
        for method in methods[contact.id]:
            if method.type == "email" and method.status == "active":
                return method.value.strip()
    return None


def mailto(db: Session, recipient: RfqRecipient) -> str | None:
    if recipient.channel != "email":
        return None
    address = email_of(db, recipient)
    if not address:
        return None
    query = f"subject={quote(recipient.subject or '')}&body={quote(recipient.message_text or '')}"
    return f"mailto:{quote(address, safe='@.+-_')}?{query}"


# ---- create / edit --------------------------------------------------------------------------

def _check_items(db: Session, case_id: uuid.UUID, item_ids: list[uuid.UUID]) -> None:
    for item_id in item_ids:
        item = items_service.get_item(db, item_id)
        if item.case_id != case_id:
            raise AppError("item_not_of_case", "The item belongs to another case", {"item_id": str(item_id)},
                           status=422)
        if item.status == "cancelled":
            raise AppError("item_cancelled", "A cancelled item cannot be sent", {"item_id": str(item_id)},
                           status=422)


def _add_recipient(db: Session, rfq: Rfq, data: dict, source: str, actor: str | None) -> RfqRecipient:
    party = parties.get_party(db, data["party_id"])
    if party.do_not_contact:
        raise AppError("do_not_contact", "This party is marked do-not-contact",
                       {"reason": party.do_not_contact_reason}, status=422)
    if data.get("contact_id") and parties.get_contact(db, data["contact_id"]).party_id != party.id:
        raise AppError("contact_not_of_party", "The contact belongs to another party", status=422)
    if db.scalar(select(RfqRecipient).where(RfqRecipient.rfq_id == rfq.id, RfqRecipient.party_id == party.id)):
        raise Conflict("duplicate_recipient", "This supplier is already a recipient of the RFQ",
                       {"party_id": str(party.id)})
    defaults = defaults_for_party(party)
    contact_id = data.get("contact_id")
    if contact_id is None:  # default: the primary active contact, else the first active one
        active = [c for c in parties.contacts_of(db, party.id) if c.status == "active"]
        contact_id = next((c.id for c in active if c.is_primary), active[0].id if active else None)
    recipient = RfqRecipient(rfq_id=rfq.id, party_id=party.id, contact_id=contact_id,
                             channel=data.get("channel") or defaults["channel"],
                             language=data.get("language") or defaults["language"], template_id=data.get("template_id"))
    db.add(recipient)
    flush(db, "duplicate_recipient")
    render_recipient(db, rfq, recipient)
    flush(db, "invalid_recipient")
    events.record(db, "rfq", rfq.id, "recipient_add", jsonable({
        "recipient_id": recipient.id, "party_id": party.id, "channel": recipient.channel,
        "language": recipient.language}), source, actor=actor)
    return recipient


def create_rfq(db: Session, data: dict, source: str, actor: str | None = None) -> Rfq:
    case = get_case(db, data["case_id"])
    if case.status != "open":
        raise Conflict("case_not_open", f"The case is {case.status}", {"status": case.status})
    item_specs = data.get("items") or []
    _check_items(db, case.id, [i["line_item_id"] for i in item_specs])
    rfq = Rfq(case_id=case.id, title=(data.get("title") or "").strip() or case.title,
              language=data.get("language") or "en", respond_by=data.get("respond_by"), notes=data.get("notes"))
    db.add(rfq)
    flush(db, "invalid_rfq")
    for spec in item_specs:
        db.add(RfqItem(rfq_id=rfq.id, line_item_id=spec["line_item_id"], quantity=spec.get("quantity"),
                       unit_code=spec.get("unit_code"), unit_text=spec.get("unit_text"), note=spec.get("note")))
    flush(db, "duplicate_item")
    events.record(db, "rfq", rfq.id, "create", jsonable({
        "case_id": case.id, "title": rfq.title, "respond_by": rfq.respond_by,
        "items": [str(i["line_item_id"]) for i in item_specs]}), source, actor=actor)
    events.record(db, "case", case.id, "rfq_add", jsonable({"rfq_id": rfq.id}), source, actor=actor)
    for recipient in data.get("recipients") or []:
        _add_recipient(db, rfq, recipient, source, actor)
    return rfq


def _draft(rfq: Rfq) -> None:
    if rfq.status in ("cancelled",):
        raise Conflict("rfq_closed", f"The RFQ is {rfq.status}", {"status": rfq.status})


def update_rfq(db: Session, rfq_id: uuid.UUID, changes: dict, source: str, actor: str | None = None) -> Rfq:
    rfq = get_rfq(db, rfq_id)
    _draft(rfq)
    item_specs = changes.pop("items", None)
    before = row_snapshot(rfq)
    for key in ("title", "language"):
        if key in changes and not changes[key]:
            changes.pop(key)
    for key, value in changes.items():
        setattr(rfq, key, value)
    if item_specs is not None:
        if any(r.status != "draft" for r in recipients_of(db, rfq.id)):
            raise Conflict("rfq_sent", "Items cannot change after the RFQ was sent; create a new RFQ")
        _check_items(db, rfq.case_id, [i["line_item_id"] for i in item_specs])
        for old, _ in rfq_items(db, rfq.id):
            db.delete(old)
        db.flush()
        for spec in item_specs:
            db.add(RfqItem(rfq_id=rfq.id, line_item_id=spec["line_item_id"], quantity=spec.get("quantity"),
                           unit_code=spec.get("unit_code"), unit_text=spec.get("unit_text"), note=spec.get("note")))
    flush(db, "invalid_rfq")
    rerender_drafts(db, rfq)
    events.record(db, "rfq", rfq.id, "update", jsonable({"before": before, "after": row_snapshot(rfq),
                                                         "items": [str(i["line_item_id"]) for i in item_specs]
                                                         if item_specs is not None else None}), source, actor=actor)
    return rfq


def rerender_drafts(db: Session, rfq: Rfq) -> None:
    for recipient in recipients_of(db, rfq.id):
        if recipient.status == "draft":
            render_recipient(db, rfq, recipient)
    db.flush()


def add_recipient(db: Session, rfq_id: uuid.UUID, data: dict, source: str, actor: str | None = None) -> RfqRecipient:
    rfq = get_rfq(db, rfq_id)
    _draft(rfq)
    return _add_recipient(db, rfq, data, source, actor)


def update_recipient(db: Session, recipient_id: uuid.UUID, changes: dict, source: str,
                     actor: str | None = None) -> RfqRecipient:
    """Draft only: channel/language/contact/template (re-renders unless text is given), subject and text."""
    recipient = get_recipient(db, recipient_id)
    rfq = get_rfq(db, recipient.rfq_id)
    if recipient.status != "draft":
        raise Conflict("recipient_sent", "The message was already sent", {"status": recipient.status})
    if changes.get("contact_id") and parties.get_contact(db, changes["contact_id"]).party_id != recipient.party_id:
        raise AppError("contact_not_of_party", "The contact belongs to another party", status=422)
    rerender = any(k in changes for k in ("channel", "language", "contact_id", "template_id")) or changes.pop(
        "rerender", False)
    before = {"channel": recipient.channel, "language": recipient.language, "subject": recipient.subject,
              "message_text": recipient.message_text}
    for key in ("channel", "language", "contact_id", "template_id"):
        if key in changes and (changes[key] is not None or key in ("contact_id", "template_id")):
            setattr(recipient, key, changes[key])
    if "template_id" in changes and changes["template_id"]:
        templates.get_template(db, changes["template_id"])
    if rerender:
        render_recipient(db, rfq, recipient)
    if "subject" in changes:
        recipient.subject = changes["subject"]
    if "message_text" in changes and changes["message_text"] is not None:
        recipient.message_text = changes["message_text"]
    flush(db, "invalid_recipient")
    events.record(db, "rfq", rfq.id, "recipient_update", jsonable({
        "recipient_id": recipient.id, "before": before,
        "after": {"channel": recipient.channel, "language": recipient.language, "subject": recipient.subject,
                  "message_text": recipient.message_text}}), source, actor=actor)
    return recipient


def remove_recipient(db: Session, recipient_id: uuid.UUID, source: str, actor: str | None = None) -> None:
    recipient = get_recipient(db, recipient_id)
    if recipient.status != "draft":
        raise Conflict("recipient_sent", "A sent message stays on record", {"status": recipient.status})
    snapshot = row_snapshot(recipient)
    db.delete(recipient)
    flush(db, "in_use")
    events.record(db, "rfq", recipient.rfq_id, "recipient_remove", {"deleted": snapshot}, source, actor=actor)
    _recompute(db, recipient.rfq_id, source, actor)


def _recompute(db: Session, rfq_id: uuid.UUID, source: str, actor: str | None, **kwargs) -> None:
    items_service.recompute_items(db, [item.id for _, item in rfq_items(db, rfq_id)], source, actor, **kwargs)


# ---- mark sent / decline / status -----------------------------------------------------------

def mark_sent(db: Session, recipient_id: uuid.UUID, data: dict, source: str, actor: str | None = None) -> RfqRecipient:
    recipient = get_recipient(db, recipient_id)
    rfq = get_rfq(db, recipient.rfq_id)
    _draft(rfq)
    if recipient.status != "draft":
        raise Conflict("recipient_sent", "Already marked as sent", {"status": recipient.status})
    if data.get("subject") is not None:
        recipient.subject = data["subject"]
    if data.get("message_text") is not None:
        recipient.message_text = data["message_text"]
    if not (recipient.message_text or "").strip():
        raise AppError("message_required", "The message text is empty", status=422)
    try:
        sent_at = ensure_aware(data.get("sent_at") or clock.now())
    except ValueError:
        raise AppError("naive_datetime", "sent_at needs a timezone", status=422) from None
    followup = data.get("followup")
    reason = data.get("no_followup_reason")
    if not followup and not reason:
        if rfq.respond_by is None:
            raise AppError("next_action_required", "Set a follow-up date (the RFQ has no respond-by date)",
                           status=422)
        followup = {"action": DEFAULT_CHASE_ACTION.get(recipient.language, DEFAULT_CHASE_ACTION["en"]),
                    "due_on": rfq.respond_by}
    channel = CHANNEL_TO_INTERACTION[recipient.channel]
    subject = recipient.subject or rfq.title
    if recipient.channel == "letter" and not recipient.subject:
        subject = f"Letter: {rfq.title}"
    interaction, created_followup = party_interactions.create_interaction(db, {
        "party_id": recipient.party_id, "contact_id": recipient.contact_id, "channel": channel, "direction": "out",
        "occurred_at": sent_at, "subject": subject, "body": recipient.message_text, "followup": followup,
        "no_followup_reason": reason}, source, actor)
    recipient.status, recipient.sent_at, recipient.interaction_id = "sent", sent_at, interaction.id
    flush(db, "invalid_recipient")
    events.record(db, "rfq", rfq.id, "sent", jsonable({
        "recipient_id": recipient.id, "party_id": recipient.party_id, "channel": recipient.channel,
        "interaction_id": interaction.id, "followup_id": created_followup.id if created_followup else None,
        "no_followup_reason": reason, "sent_at": sent_at}), source, actor=actor)
    _recompute(db, rfq.id, source, actor)
    refresh_status(db, rfq, source, actor)
    return recipient


def set_recipient_status(db: Session, recipient: RfqRecipient, status: str, source: str, actor: str | None,
                         reason: str | None = None) -> None:
    if recipient.status == status:
        return
    old = recipient.status
    recipient.status = status
    db.flush()
    events.record(db, "rfq", recipient.rfq_id, "recipient_status", jsonable({
        "recipient_id": recipient.id, "party_id": recipient.party_id, "old_status": old, "new_status": status,
        "reason": reason}), source, actor=actor)
    _recompute(db, recipient.rfq_id, source, actor)


def decline(db: Session, recipient_id: uuid.UUID, note: str | None, source: str, actor: str | None = None) -> RfqRecipient:
    recipient = get_recipient(db, recipient_id)
    if recipient.status not in ("sent", "no_response", "draft"):
        raise Conflict("recipient_closed", f"The recipient is {recipient.status}", {"status": recipient.status})
    set_recipient_status(db, recipient, "declined", source, actor, note)
    refresh_status(db, get_rfq(db, recipient.rfq_id), source, actor)
    return recipient


def refresh_status(db: Session, rfq: Rfq, source: str, actor: str | None = None) -> None:
    """draft -> sent -> answered | no_response, from the recipients (cancelled is final)."""
    if rfq.status == "cancelled":
        return
    statuses = [r.status for r in recipients_of(db, rfq.id) if r.status != "draft"]
    if not statuses:
        new = "draft"
    elif WAITING in statuses:
        new = "sent"
    elif "replied" in statuses or all(s == "declined" for s in statuses):
        new = "answered"
    else:
        new = "no_response"
    if new != rfq.status:
        old = rfq.status
        rfq.status = new
        db.flush()
        events.record(db, "rfq", rfq.id, "status", {"old_status": old, "new_status": new}, source, actor=actor)


def cancel_rfq(db: Session, rfq_id: uuid.UUID, reason: str, source: str, actor: str | None = None) -> Rfq:
    rfq = get_rfq(db, rfq_id)
    if rfq.status != "cancelled":
        old = rfq.status
        rfq.status = "cancelled"
        flush(db)
        events.record(db, "rfq", rfq.id, "cancel", {"old_status": old, "reason": reason}, source, actor=actor)
        _recompute(db, rfq.id, source, actor)
    return rfq


def delete_rfq(db: Session, rfq_id: uuid.UUID, source: str, actor: str | None = None) -> None:
    rfq = get_rfq(db, rfq_id)
    recipients = recipients_of(db, rfq.id)
    counts = {"sent": sum(1 for r in recipients if r.status != "draft"),
              "offers": db.scalar(select(func.count()).select_from(SupplierOffer).where(
                  SupplierOffer.rfq_recipient_id.in_([r.id for r in recipients] or [uuid.uuid4()]))) or 0}
    if any(counts.values()):
        raise Conflict("in_use", "The RFQ was sent or answered; cancel it instead", counts)
    snapshot = {"rfq": row_snapshot(rfq), "items": [row_snapshot(i) for i, _ in rfq_items(db, rfq.id)],
                "recipients": [row_snapshot(r) for r in recipients]}
    affected = [item.id for _, item in rfq_items(db, rfq.id)]
    for row, _ in rfq_items(db, rfq.id):
        db.delete(row)
    for row in recipients:
        db.delete(row)
    db.flush()
    db.delete(rfq)
    flush(db, "in_use")
    events.record(db, "rfq", rfq_id, "delete", {"deleted": snapshot}, source, actor=actor)
    items_service.recompute_items(db, affected, source, actor)


def expire_no_reply(db: Session, today: date | None = None, source: str = "worker") -> list[uuid.UUID]:
    """Recipients still waiting after respond_by -> no_response (status + event only; nobody is contacted)."""
    today = today or clock.today()
    rows = db.execute(select(RfqRecipient, Rfq).join(Rfq, Rfq.id == RfqRecipient.rfq_id).where(
        RfqRecipient.status == WAITING, Rfq.respond_by < today, Rfq.status != "cancelled")).all()
    changed = []
    for recipient, rfq in rows:
        set_recipient_status(db, recipient, "no_response", source, None, "respond_by passed")
        refresh_status(db, rfq, source)
        changed.append(recipient.id)
    return changed


def list_rfqs(db: Session, case_id: uuid.UUID | None = None, status: str | None = None) -> list[Rfq]:
    query = select(Rfq).order_by(Rfq.created_at.desc())
    if case_id:
        query = query.where(Rfq.case_id == case_id)
    if status:
        query = query.where(Rfq.status == status)
    return list(db.scalars(query))


# ---- export ---------------------------------------------------------------------------------

XLSX_HEADERS = {
    "en": ["No", "Item", "Description", "Quantity", "Unit", "Specs", "Packaging", "Destination", "Incoterm",
           "Required by", "HS code", "Unit price", "Currency", "MOQ", "Lead time (days)", "Remarks"],
    "fa": ["ردیف", "کالا", "شرح", "مقدار", "واحد", "مشخصات", "بسته‌بندی", "مقصد", "شرایط تحویل", "مهلت تحویل",
           "کد تعرفه", "قیمت واحد", "ارز", "حداقل سفارش", "زمان تحویل (روز)", "توضیحات"],
}


def export_xlsx(db: Session, rfq_id: uuid.UUID, language: str = "en") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    rfq = get_rfq(db, rfq_id)
    language = language if language in XLSX_HEADERS else "en"
    rows = table_rows(db, rfq.id, language)
    descriptions = [(item.description, item.hs_code) for _, item in rfq_items(db, rfq.id)]
    wb = Workbook()
    ws = wb.active
    ws.title = "RFQ"
    if language == "fa":
        ws.sheet_view.rightToLeft = True
    ws.append([rfq.title])
    ws["A1"].font = Font(bold=True)
    ws.append([])
    ws.append(XLSX_HEADERS[language])
    for cell in ws[3]:
        cell.font = Font(bold=True)
    for index, (row, (description, hs_code)) in enumerate(zip(rows, descriptions, strict=True), start=1):
        specs = "; ".join(f"{s['key']}: {s.get('value', '')}{(' ' + s['unit']) if s.get('unit') else ''}"
                          for s in row["specs"] or [])
        ws.append([index, row["name"], description or "", row["quantity"],
                   row["unit"] or "", specs, row["packaging"] or "", row["destination"] or "", row["incoterm"] or "",
                   row["required_by"] or "", hs_code or "", None, None, None, None, None])
    for column, width in zip("ABCDEFGHIJKLMNOP", (6, 36, 30, 12, 12, 40, 18, 16, 18, 14, 12, 12, 10, 12, 16, 30),
                             strict=True):
        ws.column_dimensions[column].width = width
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def recipient_out(db: Session, recipient: RfqRecipient) -> dict:
    party = parties.get_party(db, recipient.party_id)
    return {**row_snapshot(recipient), "party_name_fa": party.name_fa, "party_name_en": party.name_en,
            "party_country": party.country_iso2, "mailto": mailto(db, recipient), "email": email_of(db, recipient)}
