"""Quote outputs: print HTML, PDF (headless Edge/Chrome, no new dependency) and Word (python-docx).

Everything is rendered from the quote's own rows (locked at approval), never from live cost-sheet data, so a
regenerated file shows the same numbers. Unapproved quotes carry a "DRAFT" watermark. Files go to the document
store (hash-addressed) and are linked to the quote with role ``output``.

PDF: the print HTML (RTL/LTR by language, Vazirmatn embedded as data URIs, A4, letterhead header/footer) is written
to a temporary folder and printed by the machine's Chromium-family browser (``--headless --print-to-pdf``):
setting ``pdf.browser_path`` first, then Microsoft Edge, Chrome, Chromium. A temporary profile per run, a timeout,
networking pointed at a closed local port, cleanup afterwards. No browser -> 424 ``pdf_unavailable`` (the HTML print
view keeps working: the browser's own "Print to PDF").
"""

import base64
import html
import io
import os
import shutil
import subprocess
import tempfile
from datetime import date
from decimal import Decimal
from pathlib import Path

from sqlalchemy.orm import Session

from app.core.errors import AppError, Conflict
from app.modules.costing import service as costing
from app.modules.documents import service as documents
from app.modules.parties import service as parties
from app.modules.quotes import service
from app.modules.quotes.models import Quote
from app.modules.sourcing import coverage as sourcing_coverage
from app.modules.sourcing import items as sourcing_items
from app.modules.sourcing.jalali import to_jalali
from app.modules.system import settings

FONTS = Path(__file__).resolve().parent / "fonts"
PDF_TIMEOUT_SECONDS = 60
FA_DIGITS = str.maketrans("0123456789.,", "۰۱۲۳۴۵۶۷۸۹٫٬")
ARABIC_RANGE = ("U+0600-06FF,U+0750-077F,U+0870-088E,U+0890-0891,U+0897-08E1,U+08E3-08FF,U+200C-200E,U+2010-2011,"
                "U+204F,U+2E41,U+FB50-FDFF,U+FE70-FE74,U+FE76-FEFC")
LATIN_RANGE = ("U+0000-00FF,U+0131,U+0152-0153,U+02BB-02BC,U+02C6,U+02DA,U+02DC,U+0304,U+0308,U+0329,U+2000-206F,"
               "U+20AC,U+2122,U+2191,U+2193,U+2212,U+2215,U+FEFF,U+FFFD")

LABELS = {
    "en": {"title": "Proforma invoice", "number": "No.", "version": "Version", "date": "Date",
           "valid_until": "Valid until", "buyer": "Buyer", "attention": "Attention", "no": "#",
           "description": "Description", "quantity": "Quantity", "unit": "Unit", "unit_price": "Unit price",
           "total": "Total", "subtotal": "Subtotal", "vat": "VAT", "grand_total": "Total with VAT",
           "before_vat": "All prices are before VAT.", "with_vat": "Prices include VAT as shown.",
           "terms": "Terms", "incoterm": "Delivery (Incoterm)", "payment_terms": "Payment terms",
           "delivery_terms": "Delivery terms", "lead_time": "Lead time", "days": "days", "draft": "DRAFT",
           "draft_note": "Draft - not approved", "currency": "Currency"},
    "fa": {"title": "پیش‌فاکتور", "number": "شماره", "version": "نسخه", "date": "تاریخ",
           "valid_until": "اعتبار تا", "buyer": "خریدار", "attention": "عطف به", "no": "ردیف",
           "description": "شرح کالا", "quantity": "مقدار", "unit": "واحد", "unit_price": "قیمت واحد",
           "total": "مبلغ کل", "subtotal": "جمع", "vat": "مالیات بر ارزش افزوده", "grand_total": "جمع با مالیات",
           "before_vat": "همهٔ قیمت‌ها بدون مالیات بر ارزش افزوده است.",
           "with_vat": "قیمت‌ها با مالیات بر ارزش افزوده، مطابق جدول.",
           "terms": "شرایط", "incoterm": "شرایط تحویل (اینکوترمز)", "payment_terms": "شرایط پرداخت",
           "delivery_terms": "شرایط ارسال", "lead_time": "زمان تحویل", "days": "روز", "draft": "پیش‌نویس",
           "draft_note": "پیش‌نویس - تأیید نشده", "currency": "ارز"},
}


# ---- context (the data every output shows) --------------------------------------------------------------

def _fmt_date(value: date | None, calendar: str, digits: str) -> str:
    if value is None:
        return ""
    if calendar == "jalali":
        y, m, d = to_jalali(value)
        out = f"{y:04d}/{m:02d}/{d:02d}"
    else:
        out = value.isoformat()
    return out.translate(FA_DIGITS) if digits == "fa" else out


def _fmt_num(value: Decimal | None, decimals: int, digits: str) -> str:
    if value is None:
        return ""
    out = f"{value:,.{decimals}f}"
    return out.translate(FA_DIGITS) if digits == "fa" else out


def _fmt_qty(value: Decimal, digits: str) -> str:
    out = f"{value.normalize():,f}" if value == value.to_integral() else f"{value.normalize():f}"
    return out.translate(FA_DIGITS) if digits == "fa" else out


def build_context(db: Session, quote: Quote, calendar: str = "gregorian", digits: str = "latin") -> dict:
    lang = quote.language
    digits = digits if lang == "fa" else "latin"
    decimals = costing.currency_decimals(db, quote.currency)
    units = sourcing_coverage.unit_index(db)
    names = {i.id: i.name for i in sourcing_items.items_of(db, case_id=quote.case_id)}
    rows = []
    for row in service.items_of(db, quote.id):
        unit = units.get(row.unit_code or "")
        rows.append({"description": row.description_override or names.get(row.line_item_id, ""),
                     "note": row.note or "", "quantity": _fmt_qty(row.quantity, digits),
                     "unit": ((unit.name_fa if lang == "fa" else unit.name_en) if unit else (row.unit_text or "")),
                     "unit_price": _fmt_num(row.unit_price, decimals, digits),
                     "total": _fmt_num(row.total, decimals, digits),
                     "lead_time": f"{row.lead_time_days} {LABELS[lang]['days']}" if row.lead_time_days is not None else ""})

    def per_lang(key: str) -> str:
        return ((settings.get(db, key) or {}).get(lang) or "").strip()

    buyer = parties.get_party(db, quote.buyer_party_id) if quote.buyer_party_id else None
    contact = parties.get_contact(db, quote.contact_id) if quote.contact_id else None
    pick = (lambda fa, en: (fa or en) if lang == "fa" else (en or fa))
    logo = None
    logo_id = settings.get(db, "company.logo_document_id")
    if logo_id:
        try:
            doc = documents.get(db, logo_id)
            path = documents.file_path(doc)
            if doc.mime_type in ("image/png", "image/jpeg", "image/gif") and path.exists():
                logo = {"mime": doc.mime_type, "bytes": path.read_bytes()}
        except AppError:
            logo = None
    return {
        "lang": lang, "dir": "rtl" if lang == "fa" else "ltr", "labels": LABELS[lang],
        "draft": quote.status == "draft",
        "company": {"name": pick(settings.get(db, "company.name_fa"), settings.get(db, "company.name_en"))
                    or settings.get(db, "company.name") or "",
                    "address": per_lang("company.address"), "registration": per_lang("company.registration"),
                    "phone": settings.get(db, "company.phone") or "", "email": settings.get(db, "company.email") or "",
                    "website": settings.get(db, "company.website") or "", "footer": per_lang("company.footer"),
                    "logo": logo},
        "number": quote.number, "version": str(quote.version),
        "date": _fmt_date((quote.approved_at or quote.created_at).date(), calendar, digits),
        "valid_until": _fmt_date(quote.valid_until, calendar, digits),
        "buyer": pick(buyer.name_fa, buyer.name_en) if buyer else "",
        "contact": pick(contact.name_fa, contact.name_en) if contact else "",
        "currency": quote.currency, "rows": rows,
        "subtotal": _fmt_num(quote.subtotal, decimals, digits),
        "vat": _fmt_num(quote.vat_amount, decimals, digits),
        "total": _fmt_num(quote.total, decimals, digits),
        "with_vat": quote.price_basis == "with_vat",
        "incoterm": " ".join(x for x in (quote.incoterm, quote.incoterm_place) if x),
        "payment_terms": quote.payment_terms or "", "delivery_terms": quote.delivery_terms or "",
        "lead_time": f"{quote.lead_time_days} {LABELS[lang]['days']}" if quote.lead_time_days is not None else "",
        "terms_text": quote.terms_text or "",
    }


# ---- HTML ---------------------------------------------------------------------------------------------------

def _font_faces() -> str:
    faces = []
    for weight in (400, 700):
        for subset, rng in (("arabic", ARABIC_RANGE), ("latin", LATIN_RANGE)):
            data = base64.b64encode((FONTS / f"vazirmatn-{subset}-{weight}-normal.woff2").read_bytes()).decode()
            faces.append(f"@font-face{{font-family:'Vazirmatn';font-style:normal;font-weight:{weight};"
                         f"src:url(data:font/woff2;base64,{data}) format('woff2');unicode-range:{rng};}}")
    return "\n".join(faces)


def _e(value: str) -> str:
    return html.escape(value or "")


def _lines(value: str) -> str:
    return "<br>".join(_e(line) for line in (value or "").splitlines())


def render_html(ctx: dict) -> str:
    lb, c = ctx["labels"], ctx["company"]
    logo = ""
    if c["logo"]:
        logo = (f'<img class="logo" alt="" src="data:{c["logo"]["mime"]};base64,'
                f'{base64.b64encode(c["logo"]["bytes"]).decode()}">')
    contact_line = " · ".join(_e(x) for x in (c["phone"], c["email"], c["website"]) if x)
    rows = "".join(
        f"<tr><td class='n'>{i}</td><td><bdi>{_e(r['description'])}</bdi>"
        + (f"<div class='note'>{_e(r['note'])}</div>" if r["note"] else "")
        + (f"<div class='note'>{_e(lb['lead_time'])}: {_e(r['lead_time'])}</div>" if r["lead_time"] else "")
        + f"</td><td class='n'>{_e(r['quantity'])}</td><td>{_e(r['unit'])}</td>"
          f"<td class='n'>{_e(r['unit_price'])}</td><td class='n'>{_e(r['total'])}</td></tr>"
        for i, r in enumerate(ctx["rows"], start=1))
    totals = f"<tr><th>{_e(lb['subtotal'])}</th><td class='n'>{_e(ctx['subtotal'])} {_e(ctx['currency'])}</td></tr>"
    if ctx["with_vat"]:
        totals += (f"<tr><th>{_e(lb['vat'])}</th><td class='n'>{_e(ctx['vat'])} {_e(ctx['currency'])}</td></tr>"
                   f"<tr class='grand'><th>{_e(lb['grand_total'])}</th><td class='n'>{_e(ctx['total'])} "
                   f"{_e(ctx['currency'])}</td></tr>")
    facts = "".join(f"<tr><th>{_e(lb[key])}</th><td><bdi>{_e(ctx[key])}</bdi></td></tr>"
                    for key in ("incoterm", "payment_terms", "delivery_terms", "lead_time") if ctx[key])
    watermark = f"<div class='watermark'>{_e(lb['draft'])}</div>" if ctx["draft"] else ""
    draft_note = f"<p class='draft-note'>{_e(lb['draft_note'])}</p>" if ctx["draft"] else ""
    return f"""<!doctype html>
<html lang="{ctx['lang']}" dir="{ctx['dir']}"><head><meta charset="utf-8">
<title>{_e(lb['title'])} {_e(ctx['number'])}</title>
<style>
{_font_faces()}
@page {{ size: A4; margin: 16mm 14mm 18mm 14mm; }}
* {{ box-sizing: border-box; }}
html, body {{ margin: 0; padding: 0; }}
body {{ font-family: 'Vazirmatn', Tahoma, sans-serif; font-size: 10pt; color: #111; }}
header.letterhead {{ display: flex; align-items: center; gap: 12px; border-bottom: 2px solid #1d4ed8; padding-bottom: 6px; }}
header.letterhead .logo {{ max-height: 60px; max-width: 160px; }}
header.letterhead .name {{ font-size: 15pt; font-weight: 700; }}
header.letterhead .small {{ font-size: 8.5pt; color: #444; }}
footer.page-footer {{ position: fixed; bottom: -12mm; inset-inline: 0; font-size: 8pt; color: #555;
  border-top: 1px solid #ccc; padding-top: 3px; text-align: center; }}
h1 {{ font-size: 16pt; margin: 14px 0 6px; }}
.meta {{ display: flex; flex-wrap: wrap; gap: 4px 24px; margin-bottom: 10px; }}
.meta div span {{ color: #555; }}
table {{ width: 100%; border-collapse: collapse; }}
table.items th, table.items td {{ border: 1px solid #bbb; padding: 4px 6px; vertical-align: top; }}
table.items th {{ background: #eef2ff; text-align: start; }}
td.n {{ text-align: end; white-space: nowrap; font-variant-numeric: tabular-nums; }}
.note {{ font-size: 8.5pt; color: #555; }}
table.totals {{ width: 55%; margin-inline-start: auto; margin-top: 8px; }}
table.totals th {{ text-align: start; padding: 3px 6px; font-weight: 400; color: #333; }}
table.totals td {{ padding: 3px 6px; }}
table.totals tr.grand th, table.totals tr.grand td {{ font-weight: 700; border-top: 1px solid #333; }}
table.facts {{ margin-top: 10px; }}
table.facts th {{ text-align: start; width: 30%; padding: 2px 6px; font-weight: 400; color: #555; }}
table.facts td {{ padding: 2px 6px; }}
.basis {{ margin-top: 6px; font-weight: 700; }}
.terms {{ margin-top: 12px; white-space: normal; unicode-bidi: plaintext; }}
.watermark {{ position: fixed; top: 38%; inset-inline: 0; text-align: center; font-size: 110pt; font-weight: 700;
  color: rgba(200, 0, 0, 0.12); transform: rotate(-30deg); z-index: 0; pointer-events: none; }}
.draft-note {{ color: #b91c1c; font-weight: 700; }}
@media screen {{ body {{ max-width: 210mm; margin: 0 auto; padding: 12mm; background: #fff; }} }}
</style></head>
<body>
{watermark}
<header class="letterhead">{logo}<div><div class="name"><bdi>{_e(c['name'])}</bdi></div>
<div class="small">{_lines(c['address'])}</div><div class="small"><bdi>{contact_line}</bdi></div>
<div class="small">{_lines(c['registration'])}</div></div></header>
{draft_note}
<h1>{_e(lb['title'])}</h1>
<div class="meta">
<div><span>{_e(lb['number'])}:</span> <bdi>{_e(ctx['number'])}</bdi></div>
<div><span>{_e(lb['version'])}:</span> {_e(ctx['version'])}</div>
<div><span>{_e(lb['date'])}:</span> <bdi>{_e(ctx['date'])}</bdi></div>
<div><span>{_e(lb['valid_until'])}:</span> <bdi>{_e(ctx['valid_until'])}</bdi></div>
<div><span>{_e(lb['buyer'])}:</span> <bdi>{_e(ctx['buyer'])}</bdi></div>
{f"<div><span>{_e(lb['attention'])}:</span> <bdi>{_e(ctx['contact'])}</bdi></div>" if ctx['contact'] else ''}
<div><span>{_e(lb['currency'])}:</span> <bdi>{_e(ctx['currency'])}</bdi></div>
</div>
<table class="items"><thead><tr><th>{_e(lb['no'])}</th><th>{_e(lb['description'])}</th><th>{_e(lb['quantity'])}</th>
<th>{_e(lb['unit'])}</th><th>{_e(lb['unit_price'])} ({_e(ctx['currency'])})</th><th>{_e(lb['total'])} ({_e(ctx['currency'])})</th></tr></thead>
<tbody>{rows}</tbody></table>
<table class="totals">{totals}</table>
<p class="basis">{_e(lb['with_vat'] if ctx['with_vat'] else lb['before_vat'])}</p>
{f'<table class="facts">{facts}</table>' if facts else ''}
{f'<div class="terms"><b>{_e(lb["terms"])}</b><br>{_lines(ctx["terms_text"])}</div>' if ctx['terms_text'] else ''}
<footer class="page-footer">{_lines(c['footer'])}</footer>
</body></html>"""


# ---- PDF -----------------------------------------------------------------------------------------------------

def _candidates() -> list[str]:
    found = []
    for env in ("PROGRAMFILES(X86)", "PROGRAMFILES", "LOCALAPPDATA"):
        base = os.environ.get(env)
        if base:
            found += [str(Path(base) / "Microsoft/Edge/Application/msedge.exe"),
                      str(Path(base) / "Google/Chrome/Application/chrome.exe")]
    for name in ("msedge", "microsoft-edge", "microsoft-edge-stable", "google-chrome", "google-chrome-stable",
                 "chromium", "chromium-browser", "chrome", "chrome-headless-shell"):
        path = shutil.which(name)
        if path:
            found.append(path)
    return found


def find_browser(db: Session) -> str:
    configured = settings.get(db, "pdf.browser_path")
    if configured:
        if Path(configured).is_file():
            return configured
        raise AppError("pdf_unavailable", "The configured browser for PDF was not found; use the print view "
                       "(the browser's Print to PDF) or fix Settings -> PDF browser path",
                       {"reason": "browser_path_invalid", "path": configured}, status=424)
    for path in _candidates():
        if Path(path).is_file():
            return path
    raise AppError("pdf_unavailable", "No Edge/Chrome browser found for PDF; use the print view (the browser's "
                   "Print to PDF) or set Settings -> PDF browser path", {"reason": "browser_not_found"}, status=424)


def html_to_pdf(page: str, browser: str) -> bytes:
    with tempfile.TemporaryDirectory(prefix="boms-pdf-") as tmp:
        source, target = Path(tmp) / "quote.html", Path(tmp) / "quote.pdf"
        source.write_text(page, encoding="utf-8")
        cmd = [browser, "--headless", "--disable-gpu", "--no-first-run", "--no-default-browser-check",
               "--disable-extensions", "--disable-sync", "--disable-background-networking",
               "--disable-component-update", "--disable-default-apps", "--mute-audio",
               "--proxy-server=http://127.0.0.1:9", "--proxy-bypass-list=<-loopback>",
               f"--user-data-dir={Path(tmp) / 'profile'}", "--no-pdf-header-footer", "--print-to-pdf-no-header",
               f"--print-to-pdf={target}", source.as_uri()]
        if os.name == "posix" and hasattr(os, "geteuid") and os.geteuid() == 0:
            cmd.insert(1, "--no-sandbox")  # Chromium refuses to run as root otherwise (Linux test machines)
        try:
            subprocess.run(cmd, capture_output=True, timeout=PDF_TIMEOUT_SECONDS, check=False)
        except (OSError, subprocess.TimeoutExpired) as exc:
            raise AppError("pdf_unavailable", "The browser could not print the PDF; use the print view",
                           {"reason": type(exc).__name__}, status=424) from None
        if not target.exists() or not target.read_bytes()[:5] == b"%PDF-":
            raise AppError("pdf_unavailable", "The browser did not produce a PDF; use the print view",
                           {"reason": "no_output"}, status=424)
        return target.read_bytes()


# ---- Word ------------------------------------------------------------------------------------------------------

def render_docx(ctx: dict) -> bytes:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    rtl = ctx["dir"] == "rtl"
    font = "Tahoma" if rtl else "Calibri"  # Tahoma has Persian glyphs on every Windows; Vazirmatn may be missing
    lb, c = ctx["labels"], ctx["company"]
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    for side in ("left_margin", "right_margin"):
        setattr(section, side, Mm(14))
    section.top_margin, section.bottom_margin = Mm(16), Mm(18)
    style = doc.styles["Normal"]
    style.font.name, style.font.size = font, Pt(10)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), font)

    def bidi(paragraph):
        if rtl:
            ppr = paragraph._p.get_or_add_pPr()
            el = OxmlElement("w:bidi")
            ppr.insert(0, el)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        return paragraph

    def run(paragraph, value: str, bold: bool = False, size: float | None = None, color=None):
        r = paragraph.add_run(value)
        r.bold = bold
        if size:
            r.font.size = Pt(size)
        if color is not None:
            r.font.color.rgb = color
        if rtl:
            rpr = r._r.get_or_add_rPr()
            rpr.append(OxmlElement("w:rtl"))
            cs = OxmlElement("w:cs")
            rpr.append(cs)
        return r

    def para(container, value: str = "", **kw):
        p = bidi(container.add_paragraph())
        if value:
            run(p, value, **kw)
        return p

    def table_rtl(table):
        if rtl:
            tblpr = table._tbl.tblPr
            el = OxmlElement("w:bidiVisual")
            tblpr.append(el)

    header = section.header
    head = header.paragraphs[0]
    bidi(head)
    if c["logo"] and c["logo"]["mime"] in ("image/png", "image/jpeg", "image/gif"):
        head.add_run().add_picture(io.BytesIO(c["logo"]["bytes"]), height=Mm(14))
        head.add_run("  ")
    run(head, c["name"], bold=True, size=14)
    for value in (c["address"], " · ".join(x for x in (c["phone"], c["email"], c["website"]) if x), c["registration"]):
        if value:
            para(header, value, size=8)
    if ctx["draft"]:
        para(header, f"{lb['draft']} — {lb['draft_note']}", bold=True, size=20, color=RGBColor(0xB9, 0x1C, 0x1C))
    if c["footer"]:
        foot = section.footer.paragraphs[0]
        bidi(foot)
        foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run(foot, c["footer"], size=8)

    if ctx["draft"]:
        para(doc, lb["draft"], bold=True, size=36, color=RGBColor(0xDD, 0x99, 0x99))
    para(doc, lb["title"], bold=True, size=16)
    meta = [(lb["number"], ctx["number"]), (lb["version"], ctx["version"]), (lb["date"], ctx["date"]),
            (lb["valid_until"], ctx["valid_until"]), (lb["buyer"], ctx["buyer"])]
    if ctx["contact"]:
        meta.append((lb["attention"], ctx["contact"]))
    meta.append((lb["currency"], ctx["currency"]))
    for key, value in meta:
        p = para(doc)
        run(p, f"{key}: ", bold=True)
        run(p, value)

    headers = [lb["no"], lb["description"], lb["quantity"], lb["unit"], f"{lb['unit_price']} ({ctx['currency']})",
               f"{lb['total']} ({ctx['currency']})"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_rtl(table)
    for cell, value in zip(table.rows[0].cells, headers):
        cell.paragraphs[0].text = ""
        run(bidi(cell.paragraphs[0]), value, bold=True)
    for i, r in enumerate(ctx["rows"], start=1):
        description = r["description"] + (f"\n{r['note']}" if r["note"] else "") + \
            (f"\n{lb['lead_time']}: {r['lead_time']}" if r["lead_time"] else "")
        values = [str(i), description, r["quantity"], r["unit"], r["unit_price"], r["total"]]
        for cell, value in zip(table.add_row().cells, values):
            run(bidi(cell.paragraphs[0]), value)

    totals = [(lb["subtotal"], f"{ctx['subtotal']} {ctx['currency']}")]
    if ctx["with_vat"]:
        totals += [(lb["vat"], f"{ctx['vat']} {ctx['currency']}"), (lb["grand_total"], f"{ctx['total']} {ctx['currency']}")]
    para(doc)
    tt = doc.add_table(rows=0, cols=2)
    tt.style = "Table Grid"
    table_rtl(tt)
    for key, value in totals:
        cells = tt.add_row().cells
        run(bidi(cells[0].paragraphs[0]), key, bold=True)
        run(bidi(cells[1].paragraphs[0]), value, bold=key == lb.get("grand_total"))
    para(doc, lb["with_vat"] if ctx["with_vat"] else lb["before_vat"], bold=True)
    for key in ("incoterm", "payment_terms", "delivery_terms", "lead_time"):
        if ctx[key]:
            p = para(doc)
            run(p, f"{lb[key]}: ", bold=True)
            run(p, ctx[key])
    if ctx["terms_text"]:
        para(doc, lb["terms"], bold=True)
        for line in ctx["terms_text"].splitlines():
            para(doc, line)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ---- store and link ------------------------------------------------------------------------------------------

def generate(db: Session, quote_id, fmt: str, calendar: str, digits: str, source: str, actor: str | None) -> dict:
    """Render, store in the document store and link to the quote (role ``output``)."""
    quote = service.get_quote(db, quote_id)
    ctx = build_context(db, quote, calendar, digits)
    if fmt == "pdf":
        data, mime_ext = html_to_pdf(render_html(ctx), find_browser(db)), "pdf"
    else:
        data, mime_ext = render_docx(ctx), "docx"
    suffix = "-DRAFT" if ctx["draft"] else ""
    filename = f"{quote.number}-v{quote.version}{suffix}.{mime_ext}"
    try:
        doc = documents.upload(db, io.BytesIO(data), filename, title=filename.rsplit(".", 1)[0],
                               language=quote.language, source=source, actor=actor)
        document_id = doc.id
    except Conflict as exc:  # identical bytes were stored before: link the existing file
        if exc.code != "duplicate_document":
            raise
        document_id = exc.details["id"]
    documents.link(db, document_id, "quote", str(quote.id), "output", source=source, actor=actor)
    from app.modules.events import service as events

    events.record(db, "quote", quote.id, "output", {"format": fmt, "document_id": str(document_id),
                                                    "filename": filename, "draft": ctx["draft"]}, source, actor=actor)
    return {"document_id": str(document_id), "filename": filename, "format": fmt, "draft": ctx["draft"]}


def outputs_of(db: Session, quote_id) -> list[dict]:
    return [link for link in parties.linked_documents(db, "quote", [str(quote_id)])]
